"""Parse public SEBI portfolio workbooks by actual format and header positions.

Weights are percentage points; money is bare INR. Index-name translation and
network discovery belong to later tasks. Original source rows are never mutated.
"""
from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, replace
from decimal import Decimal, InvalidOperation

import openpyxl
import xlrd

from app.models.enums import AssetType

_XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_XLSX_MAGIC = b"PK\x03\x04"
SUM_CHECK_MIN = Decimal("90")
SUM_CHECK_MAX = Decimal("101")
# Unit detection is independent of the writer's plausibility/reconciliation gates.
FRACTION_SUM_TOLERANCE = Decimal("0.20")
_ISIN = re.compile(r"^(?:IN[A-Z0-9]{10}|IDIA[0-9]{8})$")
_HEADER_PATTERNS = {
    "instrument_name": re.compile(r"name of (the )?instrument|instrument name|company/issuer/instrument", re.I),
    "isin": re.compile(r"\bisin\b", re.I),
    "sector_raw": re.compile(r"industry|sector|rating", re.I),
    "quantity": re.compile(r"\bquantity\b", re.I),
    "market_value": re.compile(r"(?:market|mkt)[ /-]*(?:fair[ /-]*)?val(?:ue)?", re.I),
    "pct_to_nav": re.compile(r"%\s*(?:to|of)?\s*(?:net\s*)?(?:asset|nav|aum)", re.I),
}
_REQUIRED_FIELDS = ("instrument_name", "market_value", "pct_to_nav")
_SECTION_PATTERNS = [
    (AssetType.DERIVATIVE, re.compile(r"^(?:irs\b|derivatives?|(?:index / stock |stock |commodity )?(?:futures?|options?)|interest rate swaps?|exchange traded commodity derivatives)", re.I)),
    (AssetType.EQUITY_INTERNATIONAL, re.compile(r"foreign securities|overseas securit|international mutual fund|equity.*foreign", re.I)),
    (AssetType.EQUITY_DOMESTIC, re.compile(r"^arbitrage$|^equity(?: &| and| shares| instruments|$)", re.I)),
    (AssetType.DERIVATIVE, re.compile(r"^listed on commodities exchange", re.I)),
    (AssetType.OTHER, re.compile(r"^commodities(?: & commodities related)?$", re.I)),
    (AssetType.GOLD_SILVER, re.compile(r"^physical commodities$", re.I)),
    (AssetType.REIT_INVIT, re.compile(r"^(?:reits?|invits?|real estate investment trust|infrastructure investment trust|units.*(?:reit|invit|infrastructure investment trust))", re.I)),
    (AssetType.GOLD_SILVER, re.compile(r"^(?:gold|silver)(?:\s*/\s*(?:gold|silver))?(?: related| bullion|$)", re.I)),
    (AssetType.DEBT, re.compile(r"^(?:compulsory convertible debentures?|debt instruments?|corporate debt(?! market)|government securities|govt securities|central government securities|state government securities|government bonds|g-sec|bonds? & ncd|non convertible debentures?|securiti[sz]ed debt)", re.I)),
    (AssetType.CASH_EQUIVALENT, re.compile(r"^(?:money market|cash|net current (?:assets|liabilities)|nca-net current assets|net receiv|other current assets|other receiv|call, cash|treps|tri[ -]?party|reverse ?repo|cblo|repo$|short term deposits|margin money|margin on derivatives)", re.I)),
    (AssetType.OTHER, re.compile(r"^(?:others?$|units of mutual fund|mutual funds?(?: units?)?$|exchange traded funds?$|alternative investment fund|investment in mutual fund|corporate debt market development fund)", re.I)),
]
_SECTION_ONLY = re.compile(r"^(?:equity(?: &| and).*(?:related|instruments?)|debt instruments?|money market instruments?|others?|derivatives?|index / stock futures|index / stock options|mutual fund units?|exchange traded funds?|reits?/invits? instruments|listed\s*/\s*awaiting listing on stock exchanges?|government securities|non-convertible debentures\s*/\s*bonds|securiti[sz]ed debt instruments?|certificate of deposits?|commercial papers?|treasury bills?|units of an alternative investment fund(?: \(aif\))?)\s*(?:\(.*\))?$", re.I)
_TERMINAL = re.compile(r"^(?:grand[\s_]+total|total net assets|net assets)\b", re.I)
_SUBTOTAL = re.compile(r"^(?:sub\s*total|total(?:\b|:)|gross total)|\btotal$", re.I)
_NON_HOLDINGS = {"index", "scheme", "common notes", "xdo_metadata", "annexure", "f & o disclosure", "derivative"}
_AGGREGATE_CATEGORY = re.compile(
    r"^(?:units of mutual fund|units of real estate investment trusts?|units of infrastructure investment trusts?|compulsory convertible debentures?)\s*(?:\(.*\))?$", re.I,
)


class _SheetRows(list):
    # Cell formatting is a source unit cue for >100% gross FoF/ETF positions.
    percent_weight_format = False


def _uses_percent_format(number_format):
    visible = re.sub(r'"[^\"]*"|\[[^\]]*\]', "", number_format)
    return re.search(r"(?<!\\)%", visible) is not None


def _text(cell):
    return re.sub(r"\s+", " ", str("" if cell is None else cell)).strip()


def sniff_workbook_format(data: bytes) -> str:
    if data.startswith(_XLS_MAGIC):
        return "xls"
    if data.startswith(_XLSX_MAGIC):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "[Content_Types].xml" in archive.namelist() and "xl/workbook.xml" in archive.namelist():
                return "xlsx"
        raise ValueError("ZIP archive is not an OOXML workbook")
    raise ValueError(f"not a recognized workbook: first bytes {data[:8]!r}")


def market_value_multiplier(header_text: str) -> Decimal:
    lowered = header_text.lower()
    for pattern, multiplier in [(r"\bcrores?\b", "10000000"), (r"\b(?:lakhs?|lacs?)\b", "100000")]:
        if re.search(pattern, lowered):
            return Decimal(multiplier)
    for inner in re.findall(r"\(([^)]*)\)", lowered):
        remainder = re.sub(r"\b(?:rounded|rupees|inr|rs)\b|₹|[., ]", "", inner)
        if remainder:
            raise ValueError(f"unrecognized market-value unit: {header_text!r}")
    # Also reject explicit unknown non-parenthetical scales.
    if re.search(r"\bin\s+\S", lowered):
        raise ValueError(f"unrecognized market-value unit: {header_text!r}")
    return Decimal("1")


def find_header_row(rows: list[list], max_scan: int = 40) -> tuple[int, dict[str, int]] | None:
    for index, row in enumerate(rows[:max_scan]):
        columns = {}
        for col, cell in enumerate(row):
            for field, pattern in _HEADER_PATTERNS.items():
                if field not in columns and pattern.search(_text(cell)):
                    columns[field] = col
        if all(field in columns for field in _REQUIRED_FIELDS):
            return index, columns
    return None


def _to_decimal(cell):
    if cell is None or isinstance(cell, bool):
        return None
    text = str(cell).replace(",", "").replace("%", "").replace("$", "").strip().rstrip("#* ")
    try:
        value = Decimal(text)
        return value if value.is_finite() else None
    except InvalidOperation:
        return None


def _cell(row, columns, field):
    col = columns.get(field)
    return _text(row[col]) or None if col is not None and col < len(row) else None


def _label(name):
    return re.sub(r"^(?:\([a-z0-9iv]+\)|[a-ziv]+\))\s*", "", name, flags=re.I).strip()


def _section_asset_type(name, current):
    for asset_type, pattern in _SECTION_PATTERNS:
        if pattern.search(_label(name)):
            return asset_type
    return current


@dataclass(frozen=True)
class ParsedHolding:
    instrument_name: str
    isin: str | None
    sector_raw: str | None
    asset_type: AssetType
    quantity: Decimal | None
    market_value: Decimal
    pct_to_nav: Decimal


@dataclass(frozen=True)
class ParsedSheet:
    sheet_name: str
    holdings: list[ParsedHolding]
    error: str | None
    declared_total: Decimal | None = None
    scheme_title: str | None = None


def read_index_sheet(data: bytes) -> dict[str, str]:
    """Read explicit scheme tables, irrespective of the worksheet's name.

    Prefer codes which actually name a worksheet (SBI has both numeric codes and
    short codes). Conflicting names for one code are deliberately left unmapped.
    """
    fmt = sniff_workbook_format(data)
    wb = (openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
          if fmt == 'xlsx' else xlrd.open_workbook(file_contents=data))
    sheet_names = set(wb.sheetnames if fmt == 'xlsx' else wb.sheet_names())
    result, conflicts = {}, set()
    code_headers = {'schemecode', 'schemeshortcode', 'shortname', 'acronym', 'scheme', 'code', 'abbreviations', 'fundid'}
    try:
        for ws in (wb if fmt == 'xlsx' else wb.sheets()):
            rows = (list(ws.iter_rows(max_row=40, values_only=True)) if fmt == 'xlsx'
                    else [ws.row_values(i) for i in range(min(40, ws.nrows))])
            for index, row in enumerate(rows):
                labels = [re.sub(r'[^a-z0-9]', '', _text(c).lower()) for c in row]
                name_col = next((i for i, label in enumerate(labels) if label in {'schemename', 'schemenames', 'schemefullname', 'funddesc'}), None)
                code_cols = [i for i, label in enumerate(labels) if label in code_headers]
                if name_col is None or not code_cols:
                    continue
                data_rows = (ws.iter_rows(min_row=index + 2, values_only=True) if fmt == 'xlsx'
                             else (ws.row_values(i) for i in range(index + 1, ws.nrows)))
                for item in data_rows:
                    codes = [_text(item[i]) for i in code_cols if i < len(item) and _text(item[i])]
                    code = next((c for c in codes if c in sheet_names), codes[-1] if codes else None)
                    if not code or name_col >= len(item) or not _text(item[name_col]):
                        continue
                    name = str(item[name_col]).splitlines()[0].strip()
                    name = re.sub(r'\s*\((?:an?\s+)?open[- ]ended.*$', '', name, flags=re.I).strip()
                    if not name or code in conflicts:
                        continue
                    if code in result and result[code] != name:
                        result.pop(code)
                        conflicts.add(code)
                    else:
                        result[code] = name
                break
    finally:
        wb.close() if fmt == 'xlsx' else wb.release_resources()
    return result


def _row_label(row, columns):
    name = _cell(row, columns, "instrument_name")
    if name:
        return name
    # Merged headers can span several cells; total labels may sit in the rating cell.
    for cell in row[:columns["market_value"]]:
        text = _text(cell)
        if text and _to_decimal(cell) is None and not _ISIN.fullmatch(text):
            return text
    return None


def _declared_total(rows, columns):
    for row in rows:
        label = _row_label(row, columns) or ""
        if _TERMINAL.search(label):
            total = _to_decimal(_cell(row, columns, "pct_to_nav"))
            if total is None:
                raise ValueError(f"invalid declared total: {label!r}")
            return total
    return None


def _weight_decimal(cell):
    # These are the source's rounded-zero conventions, including nonzero values.
    return Decimal(0) if _text(cell) in {"^", "$", "*"} else _to_decimal(cell)


_CONTRACT = re.compile(
    r"\b(?:future|futures|option|options)\b|\b(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)[ /-]*(?:20)?\d{2}\b",
    re.I,
)


def _parse_unscaled_rows(rows: list[list], columns: dict[str, int], value_multiplier: Decimal) -> list[ParsedHolding]:
    current = AssetType.OTHER
    holdings = []
    for row_index, row in enumerate(rows):
        name = _row_label(row, columns)
        if not name:
            continue
        if _TERMINAL.search(name):
            break
        if _SUBTOTAL.search(name):
            continue
        market = _to_decimal(_cell(row, columns, "market_value"))
        pct = _weight_decimal(_cell(row, columns, "pct_to_nav"))
        isin = next((_text(c) for c in row if _ISIN.fullmatch(_text(c))), None)
        quantity = _to_decimal(_cell(row, columns, "quantity"))
        numeric_position = market is not None and pct is not None
        if re.match(r"^notes?\s*[:(]?$", name, re.I) and not isin and not numeric_position:
            break
        category = _SECTION_ONLY.fullmatch(_label(name)) or _AGGREGATE_CATEGORY.fullmatch(_label(name))
        # Explicit numeric aggregate summaries precede detailed positions in ICICI/Jio.
        # An ISIN-bearing position (even one named Government Securities) never enters here.
        aggregate = category and not isin and quantity is None and numeric_position and any(
            any(_ISIN.fullmatch(_text(c)) for c in following) or (
                not _TERMINAL.search(_row_label(following, columns) or "")
                and not _SUBTOTAL.search(_row_label(following, columns) or "")
                and _to_decimal(_cell(following, columns, "market_value")) == market
                and _weight_decimal(_cell(following, columns, "pct_to_nav")) == pct
            )
            for following in rows[row_index + 1:row_index + 6]
        )
        if aggregate:
            current = _section_asset_type(name, current)
            continue
        if category and not isin and not numeric_position:
            current = _section_asset_type(name, current)
            continue
        if market is None or pct is None:
            current = _section_asset_type(name, current)
            continue
        # Cash labels carry their own classification even inside an equity block.
        asset_type = current
        if isin is None:
            asset_type = _section_asset_type(name, current)
        sector_col = columns.get("sector_raw")
        sector = str(row[sector_col]) if sector_col is not None and sector_col < len(row) and row[sector_col] not in (None, "") else None
        if (isin is None and _CONTRACT.search(f"{name} {sector or ''}")) or (
            quantity is not None and quantity < 0
            and current in {AssetType.EQUITY_DOMESTIC, AssetType.EQUITY_INTERNATIONAL}
        ):
            asset_type = AssetType.DERIVATIVE
        elif current is not AssetType.DERIVATIVE and re.search(r"\b(?:gold|silver)\b.*\bcommodity\b", name, re.I):
            asset_type = AssetType.GOLD_SILVER
        holdings.append(ParsedHolding(name, isin, sector, asset_type,
                                      quantity, market * value_multiplier, pct))
    return holdings


def _weight_scale(holdings, percent_weight_format=False):
    if not holdings:
        return Decimal(1)
    signed = sum((h.pct_to_nav for h in holdings), Decimal(0))
    # Memo shorts may be outside NAV. They inform unit detection, never disappear
    # from the signed sum used by Write's independent reconciliation.
    without_memo_shorts = sum((h.pct_to_nav for h in holdings
                              if not (h.asset_type is AssetType.DERIVATIVE and h.pct_to_nav < 0)), Decimal(0))
    near_one = any(abs(total - Decimal(1)) <= FRACTION_SUM_TOLERANCE
                   for total in (signed, without_memo_shorts))
    bounded = max(abs(h.pct_to_nav) for h in holdings) <= Decimal(1)
    return Decimal(100) if near_one and (bounded or percent_weight_format) else Decimal(1)


def parse_sheet_rows(rows: list[list], columns: dict[str, int], value_multiplier: Decimal) -> list[ParsedHolding]:
    holdings = _parse_unscaled_rows(rows, columns, value_multiplier)
    scale = _weight_scale(holdings, getattr(rows, "percent_weight_format", False))
    return [replace(h, pct_to_nav=h.pct_to_nav * scale) for h in holdings]


def _read_xlsx_sheets(data):
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for ws in wb:
            if ws.title.lower() in _NON_HOLDINGS or ws.title.lower().startswith("derivative") or "performance" in ws.title.lower():
                continue
            rows = _SheetRows()
            blocks = False
            header_seen = False
            weight_col = None
            for cells in ws.iter_rows():
                row = [c.value for c in cells]
                rows.append(list(row))
                header = find_header_row([row])
                if header is not None:
                    header_seen = True
                    weight_col = header[1]["pct_to_nav"]
                elif weight_col is not None and weight_col < len(cells) and not any(_TERMINAL.match(_text(c)) for c in row):
                    value = _weight_decimal(row[weight_col])
                    if value is not None and value != 0:
                        rows.percent_weight_format |= _uses_percent_format(getattr(cells[weight_col], "number_format", "General"))
                blocks = blocks or any(re.fullmatch(r"SCHEME CODE\d+STARTS", _text(c)) for c in row)
                if header_seen and not blocks and any(_TERMINAL.match(_text(c)) for c in row):
                    break
            yield ws.title, rows
    finally:
        wb.close()


def _read_xls_sheets(data):
    wb = xlrd.open_workbook(file_contents=data, formatting_info=True)
    try:
        for ws in wb.sheets():
            if ws.name.lower() not in _NON_HOLDINGS and not ws.name.lower().startswith("derivative") and "performance" not in ws.name.lower():
                rows = _SheetRows(ws.row_values(i) for i in range(ws.nrows))
                header = find_header_row(rows)
                if header is not None:
                    start, columns = header
                    col = columns["pct_to_nav"]
                    for index, row in enumerate(rows[start + 1:], start + 1):
                        if _TERMINAL.search(_row_label(row, columns) or ""):
                            break
                        value = _weight_decimal(row[col]) if col < len(row) else None
                        if value is not None and value != 0:
                            key = wb.xf_list[ws.cell_xf_index(index, col)].format_key
                            rows.percent_weight_format |= _uses_percent_format(wb.format_map[key].format_str)
                yield ws.name, rows
    finally:
        wb.release_resources()


def _logical_sheets(sheets):
    """Code-delimited portfolios share one worksheet in the captured UTI file.

    Preserve code identities and the SCHEME: title for explicit name translation.
    A named portfolio total ends its positions before a defaulted-security annexure.
    """
    for name, rows in sheets:
        if not any(re.fullmatch(r"SCHEME CODE\d+STARTS", _text(c)) for row in rows for c in row):
            yield name, rows
            continue
        code = None
        block = []
        scheme = None
        ended = False
        for row in rows:
            text = next((_text(c) for c in row if _text(c)), "")
            start = re.fullmatch(r"SCHEME CODE(\d+)STARTS", text)
            finish = re.fullmatch(r"SCHEME CODE(\d+)ENDS", text)
            if start:
                code, block, scheme, ended = start.group(1), [], None, False
            elif finish and code == finish.group(1):
                yield f"{name}:{code}", block
                code = None
            elif code is not None and not ended:
                if text.startswith("SCHEME:"):
                    scheme = text.split(":", 1)[1].strip()
                if scheme and re.fullmatch(r"TOTAL\s*:\s*" + re.escape(scheme), text, re.I):
                    ended = True
                else:
                    block.append(row)


def parse_workbook(data: bytes) -> list[ParsedSheet]:
    fmt = sniff_workbook_format(data)
    sheets = _read_xlsx_sheets(data) if fmt == "xlsx" else _read_xls_sheets(data)
    parsed = []
    for name, rows in _logical_sheets(sheets):
        header = find_header_row(rows)
        if header is None:
            continue
        index, columns = header
        scheme_title = next((m.group(1).strip() for row in rows[:index] for cell in row
                             if (m := re.match(r'^SCHEME\s*:\s*(.+)$', _text(cell), re.I))), None)
        holdings = []
        total = None
        try:
            unit_text = str(rows[index][columns["market_value"]])
            if not re.search(r"lakh|lac|crore", unit_text, re.I):
                # UTI puts the scale in its title, not the MARKET-VALUE cell.
                title = " ".join(_text(c) for row in rows[:index] for c in row)
                unit_match = re.search(r"market value in (lacs?|lakhs?|crores?)", title, re.I)
                if unit_match:
                    unit_text += f" ({unit_match.group(1)})"
            multiplier = market_value_multiplier(unit_text)
            unscaled = _parse_unscaled_rows(rows[index + 1:], columns, multiplier)
            scale = _weight_scale(unscaled, getattr(rows, "percent_weight_format", False))
            holdings = [replace(h, pct_to_nav=h.pct_to_nav * scale) for h in unscaled]
            total = _declared_total(rows[index + 1:], columns)
            if total is not None:
                total_row = next(r for r in rows[index + 1:] if _TERMINAL.search(_row_label(r, columns) or ""))
                explicit_percent = "%" in (_cell(total_row, columns, "pct_to_nav") or "")
                declared_scale = Decimal(1) if explicit_percent or abs(total) > Decimal("1.5") else Decimal(100)
                total *= declared_scale
                if not explicit_percent and declared_scale != scale and any(h.pct_to_nav for h in unscaled):
                    raise ValueError("declared total scale disagrees with row weights")
            parsed.append(ParsedSheet(name, holdings, None, total, scheme_title))
        except Exception as exc:
            parsed.append(ParsedSheet(name, holdings, f"{type(exc).__name__}: {exc}", total, scheme_title))
    return parsed


def extract_zip_members(data: bytes) -> list[tuple[str, bytes]]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return [(n, archive.read(n)) for n in archive.namelist()
                if n.lower().endswith((".xls", ".xlsx")) and not n.endswith("/")]
