"""Batch-only Locate/Fetch/Parse/Normalize/Write orchestration.

The synchronous Session queries/flushes inside async functions are intentional for
the scheduled job; commits use commit_off_loop. Never call this from an API route.
Each AMC commits separately, so a later failure cannot undo an earlier success.
Discovery keeps lightweight references; the refresh iterator downloads, parses and
writes one file at a time. Lateness alerts are gated; file errors remain isolated.
"""
from __future__ import annotations

import html
import ast
import json
import logging
import re
import time
import urllib.parse
import zipfile
import io
import uuid
from contextvars import ContextVar
from collections.abc import AsyncIterator
from dataclasses import dataclass
from datetime import date, timedelta

import httpx
import openpyxl
import xlrd
from sqlalchemy.orm import Session

from app.db.session import commit_off_loop
from app.models.enums import DisclosureBatchStatus
from app.models.reference import DisclosureBatch, Scheme
from app.services.analytics.amfi_factsheet_client import build_families
from app.services.analytics.portfolio_disclosure_matching import match_scheme_name
from app.services.analytics.portfolio_disclosure_parsing import (
    extract_zip_members, parse_workbook, sniff_workbook_format, read_index_sheet,
)
from app.services.analytics.portfolio_disclosure_resolvers import (
    AMC_DISCLOSURE_RESOLVERS, DisclosureFileShape, DisclosureResolverEntry, DisclosureResolverKind,
)
from app.services.analytics.portfolio_disclosure_writer import write_disclosure_batch
from app.services.analytics.portfolio_disclosure_retention import archive_failed_file

logger = logging.getLogger(__name__)
_HTTP_TIMEOUT = 90.0
_DISCOVERY_ONLY = ContextVar('portfolio_disclosure_discovery_only', default=False)
_RUN_TODAY = ContextVar('portfolio_disclosure_run_today', default=None)
_ALERT_DAYS = (10, 20)
_HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"}
_MAX_WORKBOOK_BYTES = 60 * 1024 * 1024
_ANCHOR_LINK_RE = re.compile(r'<a\b[^>]*\bhref=["\']([^"\']+\.(?:xlsx?|zip)(?:[?#][^"\']*)?)["\'][^>]*>(.*?)</a>', re.I | re.S)
_BARE_LINK_RE = re.compile(r'["\']([^"\'<>\r\n]+\.(?:xlsx?|zip)(?:[?#][^"\'<>\r\n]*)?)["\']', re.I)
_MONTH_PATTERN = r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)"
_MONTHS = {m: i for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}
_SEPARATOR = r"[\s_.–-]+"
_DAY_DATE_RE = re.compile(rf"(?<!\d)(\d{{1,2}})(?:st|nd|rd|th)?{_SEPARATOR}({_MONTH_PATTERN}|\d{{2}}){_SEPARATOR}(\d{{4}}|\d{{2}})(?!\d)", re.I)
_COMPACT_DATE_RE = re.compile(r"(?<!\d)(\d{2})(\d{2})(\d{4})(?!\d)")
_MONTH_YEAR_RE = re.compile(rf"(?<![A-Za-z])({_MONTH_PATTERN})[\s_.–-]*(\d{{4}}|\d{{2}})(?!\d)", re.I)


@dataclass(frozen=True)
class LocatedFile:
    data: bytes | None  # None is a lightweight discovery reference, never parsed.
    source_url: str
    as_of_date: date
    scheme_hint: str | None = None


@dataclass(frozen=True)
class PortfolioDisclosureRefreshResult:
    success: bool
    amcs_processed: int = 0
    amcs_failed: int = 0
    amcs_not_onboarded: int = 0
    schemes_written: int = 0
    seconds: float = 0.0  # elapsed time only; never a financial value


def _filename_as_of_date(url: str) -> date | None:
    filename = urllib.parse.unquote(urllib.parse.urlsplit(url).path.rsplit("/", 1)[-1])
    match = _DAY_DATE_RE.search(filename) or _COMPACT_DATE_RE.search(filename)
    if match:
        day, month, year = match.groups()
        month_number = int(month) if month.isdigit() else _MONTHS[month[:3].lower()]
        year_number = int(year) + (2000 if len(year) == 2 else 0)
        try:
            return date(year_number, month_number, int(day))
        except ValueError:
            return None
    match = _MONTH_YEAR_RE.search(filename)
    if match:
        month, year = match.groups()
        year_number = int(year) + (2000 if len(year) == 2 else 0)
        try:
            next_month = date(year_number, _MONTHS[month[:3].lower()], 28) + timedelta(days=4)
            return next_month - timedelta(days=next_month.day)
        except ValueError:
            return None
    return None


def _extract_links(html_text: str) -> list[tuple[str, str]]:
    # Decode only JSON URL/quote escapes, preserving percent-encoded paths and text.
    text = re.sub(r"\\+u0026", "&", html_text, flags=re.I)
    text = re.sub(r"\\+([/\"'])", r"\1", text)
    text = html.unescape(text)
    found = {}
    for href, label in _ANCHOR_LINK_RE.findall(text):
        found[href] = " ".join(re.sub(r"<[^>]+>", " ", label).split())
    for href in _BARE_LINK_RE.findall(text):
        found.setdefault(href, "")
    return list(found.items())


def _clean_label(label: str) -> str:
    label = html.unescape(" ".join(label.split())).strip()
    # Canara prefix discovery has no anchor text or Index sheet. Its captured
    # filename explicitly names the scheme between the leading code and date.
    decoded = urllib.parse.unquote(label)
    canara = re.fullmatch(
        r"[A-Za-z]{1,4}[-_\s–]+(Canara-Robeco-.*?-Fund)[-_\s–]+.*\.xlsx?",
        decoded, re.I,
    )
    if canara:
        return canara.group(1).replace("-", " ")
    # The captured per-scheme anchor text includes a document descriptor.
    return re.sub(r"\s*[-–:]\s*Monthly Portfolio(?:\s.*)?$", "", label, flags=re.I).strip()


_WORKBOOK_URL = re.compile(r"\.(?:xlsx?|zip)(?:$|[?#])", re.I)
_WORKBOOK_MAGICS = (b"PK\x03\x04", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")


def _looks_like_workbook(data: bytes) -> bool:
    return data.startswith(_WORKBOOK_MAGICS)


def _candidate_allowed(entry, url, label=""):
    text = urllib.parse.unquote(f"{url} {label}")
    return (
        (not entry.link_pattern or re.search(entry.link_pattern, text, re.I) is not None)
        and (not entry.exclude_pattern or re.search(entry.exclude_pattern, text, re.I) is None)
    )


async def _static_regex_candidates(client: httpx.AsyncClient, entry: DisclosureResolverEntry) -> list[tuple[str, str]]:
    resp = await client.get(entry.landing_url)
    resp.raise_for_status()
    return [(urllib.parse.urljoin(str(resp.url), href), label)
            for href, label in _extract_links(resp.text) if _candidate_allowed(entry, href, label)]


async def _download_workbook(client: httpx.AsyncClient, url: str) -> bytes | None:
    async with client.stream("GET", url) as resp:
        resp.raise_for_status()
        body = bytearray()
        async for chunk in resp.aiter_bytes():
            body.extend(chunk)
            if len(body) > _MAX_WORKBOOK_BYTES:
                return None
    # HTTP 200 is not proof of a file (Mirae-style HTML under a .xlsx name): check the magic bytes.
    return bytes(body) if _looks_like_workbook(bytes(body[:8])) else None


async def _fetch_candidates(client: httpx.AsyncClient, candidates: list[tuple[str, str]], *, keep_all: bool) -> list[LocatedFile]:
    dated = sorted(((url, label, as_of) for url, label in candidates
                    if (as_of := _filename_as_of_date(url)) is not None),
                   key=lambda item: item[2], reverse=True)
    if keep_all and dated:
        # Prefer the target month; otherwise the newest available non-future month.
        # Keep every date in that month so a stray later date cannot drop its peers.
        target = _locate_month_end()
        eligible = [item for item in dated if item[2] <= target]
        month = (target.year, target.month)
        if not any((item[2].year, item[2].month) == month for item in eligible) and eligible:
            month = (eligible[0][2].year, eligible[0][2].month)
        dated = [item for item in eligible if (item[2].year, item[2].month) == month]
    results = []
    seen = set()
    for url, label, as_of in dated:
        if url in seen:
            continue
        seen.add(url)
        if _DISCOVERY_ONLY.get():
            results.append(LocatedFile(None, url, as_of, _clean_label(label) or None if keep_all else None))
            continue
        try:
            data = await _download_workbook(client, url)
        except httpx.HTTPError:
            continue
        if data:
            results.append(LocatedFile(data, url, as_of, _clean_label(label) or None if keep_all else None))
            if not keep_all:
                break
    return results


def _slugify_scheme_name(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def _latest_month_end(today: date | None = None) -> date:
    return (today or date.today()).replace(day=1) - timedelta(days=1)


def _locate_month_end() -> date:
    today = _RUN_TODAY.get()
    return _latest_month_end(today) if today is not None else _latest_month_end()


async def _locate_deterministic_template(client: httpx.AsyncClient, entry: DisclosureResolverEntry, scheme_names: list[str]) -> list[LocatedFile]:
    as_of = _locate_month_end()
    results = []
    seen = set()
    for name in scheme_names:
        url = entry.url_template.format(
            scheme_slug=_slugify_scheme_name(name), dd=f"{as_of.day:02d}", mm=f"{as_of.month:02d}",
            yyyy=as_of.year, month_name_lower=as_of.strftime("%B").lower(),
            month_name=as_of.strftime("%B"), yy=f"{as_of.year % 100:02d}",
            yy_plus_1=f"{(as_of.year + 1) % 100:02d}",
        )
        if url in seen:
            continue
        seen.add(url)
        if _DISCOVERY_ONLY.get():
            results.append(LocatedFile(None, url, as_of, name if '{scheme_slug}' in entry.url_template else None))
            continue
        try:
            data = await _download_workbook(client, url)
        except httpx.HTTPError:
            continue
        if data:
            hint = name if "{scheme_slug}" in entry.url_template else None
            results.append(LocatedFile(data, url, as_of, hint))
    return results


def _is_locate_implemented(entry: DisclosureResolverEntry) -> bool:
    if entry.kind is DisclosureResolverKind.MANUAL:
        return False
    if entry.custom_locator:
        return entry.custom_locator in _CUSTOM_LOCATORS
    if entry.kind is DisclosureResolverKind.JSON_API:
        # Unconfigured page-payload rows must remain counted as not onboarded.
        return bool(entry.endpoint_url and entry.response_json_path is not None and entry.documents is not None)
    return (entry.kind in {DisclosureResolverKind.STATIC_REGEX, DisclosureResolverKind.WAYBACK_LOCATE}
            or entry.file_shape is DisclosureFileShape.DETERMINISTIC_TEMPLATE)


async def _locate_and_fetch(client: httpx.AsyncClient, entry: DisclosureResolverEntry, scheme_rows: list[Scheme]) -> list[LocatedFile]:
    if entry.custom_locator:
        return await _CUSTOM_LOCATORS[entry.custom_locator](client, entry, scheme_rows)
    if entry.kind is DisclosureResolverKind.JSON_API and _is_locate_implemented(entry):
        return await _locate_json_api(client, entry, _locate_month_end())
    if entry.kind is DisclosureResolverKind.WAYBACK_LOCATE:
        return await _locate_wayback(client, entry)  # implemented in Task 7
    if entry.file_shape is DisclosureFileShape.DETERMINISTIC_TEMPLATE:
        return await _locate_deterministic_template(client, entry, [s.base_name or s.name for s in scheme_rows])
    if entry.kind is not DisclosureResolverKind.STATIC_REGEX:
        raise NotImplementedError(f"{entry.kind} Locate requires its own onboarding task")
    return await _fetch_candidates(client, await _static_regex_candidates(client, entry),
                                   keep_all=entry.file_shape is DisclosureFileShape.MULTI_FILE_SINGLE_SCHEME)


def _request_values(value, as_of: date):
    """Fill only named placeholders; regex quantifier braces remain literal."""
    fy = as_of.year if as_of.month >= 4 else as_of.year - 1
    fields = {'yyyy': str(as_of.year), 'year': str(as_of.year), 'yy': f'{as_of.year % 100:02d}',
              'yy_plus_1': f'{(as_of.year + 1) % 100:02d}', 'mm': f'{as_of.month:02d}',
              'month': str(as_of.month), 'month_name': as_of.strftime('%B'),
              'financial_year': f'{fy}-{fy + 1}'}
    if isinstance(value, str):
        for key, replacement in fields.items():
            value = value.replace('{' + key + '}', replacement)
    elif isinstance(value, dict):
        return {key: _request_values(item, as_of) for key, item in value.items()}
    elif isinstance(value, (list, tuple)):
        return [_request_values(item, as_of) for item in value]
    return value


def _response_json(response):
    response.raise_for_status()
    if response.status_code == 204 or not response.content:
        return {}  # "no documents" (Franklin answers 204 for funds with none), not an error
    if not response.headers.get('Content-Type', '').lower().startswith('application/json'):
        raise ValueError('BLOCKED: expected JSON response, received non-JSON content type')
    return response.json()


async def _json_request(client, entry, as_of, *, headers=None, body=None, params=None):
    kwargs = {'params': _request_values(entry.request_params if params is None else params, as_of),
              'headers': {**_HEADERS, **(headers or {})}}
    if entry.http_method.upper() == 'POST':
        kwargs['json'] = _request_values(entry.request_body if body is None else body, as_of)
    response = await client.request(entry.http_method, entry.endpoint_url, **kwargs)
    return _response_json(response)


def _disclosed_date(text: str | None) -> date | None:
    """Parse a disclosure title, never an upload/reference/published timestamp."""
    text = html.unescape(str(text or '')).strip()
    text = re.split(r'\(?\bUploaded on\b', text, maxsplit=1, flags=re.I)[0].strip()
    # Franklin's month-first title precedes its parenthetical upload date.
    month_first = re.search(rf'({_MONTH_PATTERN})\s+(\d{{1,2}})(?:st|nd|rd|th)?\s*,\s*(\d{{4}})', text, re.I)
    if month_first:
        month, day, year = month_first.groups()
        try:
            return date(int(year), _MONTHS[month[:3].lower()], int(day))
        except ValueError:
            return None
    # WhiteOak writes "30th September2026" without a space before the year.
    text = re.sub(rf'({_MONTH_PATTERN})(?=\d{{4}})', r'\1 ', text, flags=re.I)
    text = text.replace(',', ' ')
    return _filename_as_of_date('https://disclosure.invalid/' + urllib.parse.quote(text, safe=''))


def _file_url(url: str, base: str, rewrite_host=None):
    joined = urllib.parse.urljoin(base, html.unescape(url).strip())
    parsed = urllib.parse.urlsplit(joined)
    if parsed.scheme not in {'http', 'https'}:
        return None
    if rewrite_host and parsed.netloc.lower() == rewrite_host[0].lower():
        joined = parsed._replace(netloc=rewrite_host[1]).geturl()
    return str(httpx.URL(joined))


async def _fetch_dated_records(client, records, as_of, *, keep_all=True):
    """Source metadata can date opaque URLs; an undated source never gets a default."""
    results, seen = [], set()
    for url, disclosed, hint in records:
        if not url or disclosed != as_of or url in seen or not _WORKBOOK_URL.search(url):
            continue
        seen.add(url)
        if _DISCOVERY_ONLY.get():
            results.append(LocatedFile(None, url, disclosed, hint))
            continue
        try:
            data = await _download_workbook(client, url)
        except httpx.HTTPError:
            continue
        if data:
            results.append(LocatedFile(data, url, disclosed, hint))
            if not keep_all:
                break
    return results


async def _locate_json_api(client, entry, as_of) -> list[LocatedFile]:
    rows = await _json_request(client, entry, as_of)
    for key in (entry.response_json_path or '').split('.'):
        if key:
            rows = rows.get(key, []) if isinstance(rows, dict) else []
    records = []
    for row in rows if isinstance(rows, list) else []:
        if isinstance(row, str):
            url, label = row, ''
        elif isinstance(row, dict):
            url = row.get('doc') or row.get('pdfUrl') or row.get('url')
            label = row.get('ResourceLink') or row.get('name') or ''
        else:
            continue
        if not isinstance(url, str):
            continue
        if entry.documents and not any(re.search(_request_values(pattern, as_of), f'{label} {url}', re.I)
                                       for pattern in entry.documents):
            continue
        url = _file_url(url, entry.endpoint_url, entry.rewrite_host)
        disclosed = _disclosed_date(label) or (_filename_as_of_date(url) if url else None)
        records.append((url, disclosed, None))
    return await _fetch_dated_records(client, records, as_of,
                                    keep_all=entry.file_shape is DisclosureFileShape.MULTI_FILE_SINGLE_SCHEME)


async def _locate_franklin(client, entry, scheme_rows):
    as_of = _locate_month_end()
    sitemap = await client.get(entry.landing_url)
    sitemap.raise_for_status()
    ids = dict.fromkeys(re.findall(r'/fund-details/fund-overview/(\d+)', sitemap.text))
    records = []
    for fund_id in ids:
        params = {**(entry.request_params or {}), 'frkOneTISNumber': fund_id}
        payload = await _json_request(client, entry, as_of, params=params)
        for row in payload.get('document', []):
            if row.get('contentGrouping') != 'Dis-Portfolio-Monthly':
                continue
            url = _file_url(row.get('downloadUrl', ''), entry.endpoint_url)
            filename = urllib.parse.unquote(urllib.parse.urlsplit(url or '').path.rsplit('/', 1)[-1])
            hint = re.sub(rf'_{_MONTH_PATTERN}_\d{{2,4}}_Monthly_Portfolio\.xlsx?$', '', filename, flags=re.I)
            records.append((url, _disclosed_date(row.get('dctermsTitle')), hint.replace('_', ' ')))
    return await _fetch_dated_records(client, records, as_of)


async def _locate_axis(client, entry, scheme_rows):
    as_of = _locate_month_end()
    payload = await _json_request(client, entry, as_of)
    token_data = payload.get('data', {})
    if isinstance(token_data, str):
        token_data = ast.literal_eval(token_data)
    token = token_data.get('token', '').strip()
    if not token:
        raise ValueError('Axis token missing')
    authorization = token if token.lower().startswith('bearer ') else 'Bearer ' + token
    response = await client.post(urllib.parse.urljoin(entry.endpoint_url, '/cms/get-scheme-documents'),
                                 headers={'Authorization': authorization}, json={
                                     'sdType': 'yearMonthSchemeDocs', 'sdID': 'sdMonthSchemePortfolio',
                                     'year': str(as_of.year), 'month': as_of.strftime('%B'), 'schemeCode': 'Consolidated'})
    rows = _response_json(response).get('data', {}).get('documentList', [])
    records = [(_file_url(r.get('docuementURL', ''), entry.endpoint_url), _disclosed_date(r.get('documentName')), None)
               for r in rows if str(r.get('documentName', '')).lower().startswith('monthly portfolio')]
    return await _fetch_dated_records(client, records, as_of, keep_all=False)


def _html_records(text, base, *, hints=False):
    return [(_file_url(url, base), _disclosed_date(label) or _filename_as_of_date(url.strip()),
             _clean_label(label) if hints else None) for url, label in _extract_links(text)]


def _options(text, select_name=None):
    if select_name:
        found = re.search(r'<select\b[^>]*name=["\']' + re.escape(select_name) + r'["\'][^>]*>(.*?)</select>', text, re.I | re.S)
        text = found.group(1) if found else ''
    return [(value.strip(), ' '.join(re.sub(r'<[^>]+>', '', html.unescape(label)).split()))
            for value, label in re.findall(r'<option\b[^>]*value=["\']([^"\']*)["\'][^>]*>(.*?)</option>', text, re.I | re.S)]


async def _locate_lic(client, entry, scheme_rows):
    as_of = _locate_month_end()
    headers = {'X-Requested-With': 'XMLHttpRequest'}
    year = await client.post(entry.endpoint_url, data=entry.request_body, headers=headers)
    year.raise_for_status()
    if str(as_of.year) not in {v for v, label in _options(year.text)}:
        return []
    month = await client.post(entry.endpoint_url, data={'year': str(as_of.year), 'id': '639', 'filter': 'month'}, headers=headers)
    month.raise_for_status()
    values = {v for v, label in _options(month.text)}
    month_value = next((v for v in values if v.isdigit() and int(v) == as_of.month), None)
    if month_value is None:
        return []
    response = await client.post(urllib.parse.urljoin(entry.endpoint_url, '/downloads/consolidated-portfolio-files'),
                                 data={'id': '639', 'month': month_value, 'year': str(as_of.year)}, headers=headers)
    response.raise_for_status()
    # Both consolidated workbooks cover different schemes; neither is a scheme hint.
    return await _fetch_dated_records(client, _html_records(response.text, str(response.url)), as_of)


async def _locate_sbi(client, entry, scheme_rows):
    as_of = _locate_month_end()
    response = await client.post(entry.endpoint_url, json=_request_values(entry.request_body, as_of),
                                 headers={'X-Requested-With': 'XMLHttpRequest', 'Content-Type': 'application/json;charset=UTF-8'})
    response.raise_for_status()
    records = [r for r in _html_records(response.text, str(response.url)) if r[0] and 'all-schemes-monthly-portfolio' in r[0].lower()]
    return await _fetch_dated_records(client, records, as_of, keep_all=False)


async def _locate_trust(client, entry, scheme_rows):
    as_of = _locate_month_end()
    payload = await _json_request(client, entry, as_of)
    records = [(_file_url(r.get('fileurl', ''), entry.endpoint_url), _disclosed_date(r.get('title')), None)
               for r in payload.get('resultSetArray', [])]
    return await _fetch_dated_records(client, records, as_of, keep_all=False)


async def _locate_mirae(client, entry, scheme_rows):
    as_of = _locate_month_end()
    records, page = [], 1
    while True:
        body = _request_values(entry.request_body, as_of)
        body['request']['pgno'] = page
        payload = await _json_request(client, entry, as_of, body=body, headers={'X-Requested-With': 'XMLHttpRequest'})
        rows = payload.get('Data', [])
        for row in rows:
            title = row.get('Title', '')
            hint = title.split(' for ', 1)[1].strip() if ' for ' in title else None
            records.append((_file_url(row.get('URL', ''), entry.endpoint_url), _disclosed_date(title), hint))
        if not rows or page * body['request']['pgsize'] >= int(payload.get('DataCount', len(rows))):
            break
        page += 1
    return await _fetch_dated_records(client, records, as_of)


async def _locate_boi(client, entry, scheme_rows):
    as_of = _locate_month_end()
    payload = await _json_request(client, entry, as_of)
    data = payload.get('d', {})
    if isinstance(data, str):
        data = json.loads(data)
    records = [(_file_url(r.get('FolderUrl', ''), entry.endpoint_url), _disclosed_date(r.get('DocName')), None)
               for r in data.get('Documents', [])]
    return await _fetch_dated_records(client, records, as_of, keep_all=False)


async def _locate_taurus(client, entry, scheme_rows):
    completed = _latest_month_end()
    year, month = completed.year, 9 if completed.month >= 9 else 3
    if completed.month < 3:
        year, month = year - 1, 9
    as_of = date(year, month + 1, 1) - timedelta(days=1)
    response = await client.get(entry.endpoint_url)
    response.raise_for_status()
    year_id = next((v for v, label in _options(response.text, 'field_hy_portfolio_year_target_id') if label == str(year)), None)
    month_id = next((v for v, label in _options(response.text, 'field_hy_portfolio_month_target_id') if label.lower() == as_of.strftime('%B').lower()), None)
    if year_id is None or month_id is None:
        return []
    response = await client.get(entry.endpoint_url, params={'field_hy_portfolio_year_target_id': year_id,
                                                            'field_hy_portfolio_month_target_id': month_id})
    response.raise_for_status()
    records = []
    for url, label in _extract_links(response.text):
        url = _file_url(url, str(response.url))
        filename = urllib.parse.unquote(urllib.parse.urlsplit(url or '').path.rsplit('/', 1)[-1])
        hint = re.sub(r'_Half_Yearly_Portfolio.*$', '', filename, flags=re.I).replace('_', ' ')
        records.append((url, _filename_as_of_date(url) if url else None, hint))
    return await _fetch_dated_records(client, records, as_of)


async def _locate_absl(client, entry, scheme_rows):
    return await _locate_json_api(client, entry, _locate_month_end())


async def _locate_navi(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url)
    page.raise_for_status()
    nonce = re.search(r'"rest_url"\s*:\s*"[^"]*"\s*,\s*"nonce"\s*:\s*"([a-z0-9]+)"', page.text)
    if nonce is None:
        raise ValueError('Navi WP nonce missing')
    response = await client.post(entry.endpoint_url, data=_request_values(entry.request_body, as_of), headers={'WP-NONCE': nonce.group(1)})
    rows = _response_json(response).get('data', [])
    records = []
    for row in rows:
        title = html.unescape(row.get('title', ''))
        hint = re.split(r'\s+\d{1,2}(?:st|nd|rd|th)\b', title, maxsplit=1, flags=re.I)[0].strip()
        records.append((_file_url(row.get('url', ''), entry.endpoint_url), _disclosed_date(title), hint))
    return await _fetch_dated_records(client, records, as_of)


async def _locate_pgim(client, entry, scheme_rows):
    as_of = _locate_month_end()
    payload = await _json_request(client, entry, as_of)
    records = []
    for tab in payload.get('data', []):
        for row in tab.get('content', []):
            title = row.get('title', '')
            hint = re.sub(rf'\s+{_MONTH_PATTERN}\s+\d{{4}}$', '', title, flags=re.I).strip()
            records.append((_file_url(row.get('pdfPath', ''), entry.endpoint_url), _disclosed_date(row.get('dateMonthYear')), hint))
    return await _fetch_dated_records(client, records, as_of)


async def _locate_whiteoak(client, entry, scheme_rows):
    as_of = _locate_month_end()
    records, page = [], 1
    while True:
        params = {**entry.request_params, 'pagination[page]': page}
        payload = await _json_request(client, entry, as_of, params=params)
        for row in payload.get('data', []):
            attrs = row.get('attributes', {})
            if attrs.get('period') != 'Monthly':
                continue
            file = (attrs.get('doc_file') or {}).get('data') or {}
            records.append((_file_url(file.get('attributes', {}).get('url', ''), entry.endpoint_url),
                            _disclosed_date(attrs.get('doc_name')), attrs.get('scheme_name')))
        if page >= int(payload.get('meta', {}).get('pagination', {}).get('pageCount', 1)):
            break
        page += 1
    return await _fetch_dated_records(client, records, as_of)


async def _locate_icici(client, entry, scheme_rows):
    as_of = _locate_month_end()
    headers = {'requestapiid': str(uuid.uuid4()), 'sourceurl': 'DOWNLOADS', 'env': 'api',
               'Origin': 'https://www.icicipruamc.com', 'Referer': entry.landing_url,
               'Accept': 'application/json'}
    payload = await _json_request(client, entry, as_of, headers=headers)
    rows = payload.get('success', {}).get('data', {}).get('files', [])
    records = []
    for row in rows:
        path = row.get('url', '')
        if not path.startswith('/downloads/'):
            continue
        url = _file_url('/blob' + path, 'https://www.icicipruamc.com')
        records.append((url, _disclosed_date(row.get('title', {}).get('text')), None))
    return await _fetch_dated_records(client, records, as_of, keep_all=False)


async def _locate_uti(client, entry, scheme_rows):
    # No blanket scheme hint: EXPOSURE:<code> is not a tested scheme name.
    return await _locate_json_api(client, entry, _locate_month_end())


def _walk_dicts(value):
    if isinstance(value, dict):
        yield value
        for child in value.values():
            yield from _walk_dicts(child)
    elif isinstance(value, list):
        for child in value:
            yield from _walk_dicts(child)


def _json_fragments(text):
    """Decode complete JSON values in script/flight text without executing JS."""
    decoder, offset = json.JSONDecoder(), 0
    while found := re.search(r'[\[{]', text[offset:]):
        start = offset + found.start()
        try:
            value, end = decoder.raw_decode(text, start)
        except json.JSONDecodeError:
            offset = start + 1
        else:
            yield value
            offset = end


def _page_objects(text):
    def visit(value):
        if isinstance(value, dict):
            yield value
            for child in value.values():
                yield from visit(child)
        elif isinstance(value, list):
            for child in value:
                yield from visit(child)
        elif isinstance(value, str):
            for decoded in _json_fragments(value):
                yield from visit(decoded)
    for source in (text, re.sub(r'\\+([/"\'])', r'\1', text)):
        for value in _json_fragments(source):
            yield from visit(value)


def _workbook_title_rows(data, count=12):
    """Read disclosed title cells only; never infer a scheme from its short code."""
    try:
        fmt = sniff_workbook_format(data)
    except (ValueError, zipfile.BadZipFile):
        return []
    if fmt == 'xlsx':
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            return list(wb.worksheets[0].iter_rows(max_row=count, values_only=True))
        finally:
            wb.close()
    wb = xlrd.open_workbook(file_contents=data)
    try:
        ws = wb.sheet_by_index(0)
        return [ws.row_values(i) for i in range(min(count, ws.nrows))]
    finally:
        wb.release_resources()


async def _locate_jio(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url); page.raise_for_status()
    scripts = re.findall(r'<script[^>]+src=["\']([^"\']+)["\']', page.text)
    action = None
    for script in scripts:
        if 'statutory-disclosure/' not in script or '/page-' not in script:
            continue
        code = await client.get(urllib.parse.urljoin(str(page.url), html.unescape(script)))
        code.raise_for_status()
        action = re.search(r'createServerReference\)\("([^"\s]+)"[^;()]+?"getDisclosureL3Data"', code.text)
        if action:
            break
    if not action:
        return []
    fy = as_of.year if as_of.month >= 4 else as_of.year - 1
    response = await client.post(entry.endpoint_url, headers={'Next-Action':action.group(1),
        'Content-Type':'text/plain;charset=UTF-8','Accept':'text/x-component'}, content=json.dumps([
            'monthly-portfolio-disclosure', {'year':f'FI{fy}-{fy+1}','month':as_of.strftime('%B'),'date':'$undefined'}, 'MF']))
    response.raise_for_status()
    payload = next((json.loads(line[2:]) for line in response.text.splitlines() if line.startswith('1:')), {})
    rows = payload.get('data', []) if isinstance(payload, dict) else []
    all_schemes = [r for r in rows if re.match(r'Jio\s*BlackRock Mutual Fund\b', r.get('title',''), re.I)]
    records = []
    for row in all_schemes or rows:
        title = row.get('title','')
        hint = None if all_schemes else re.split(r'-Monthly-Portfolio', title, flags=re.I)[0].strip()
        records.append((_file_url((row.get('file') or {}).get('url',''), entry.endpoint_url), _disclosed_date(title), hint))
    return await _fetch_dated_records(client, records, as_of, keep_all=not bool(all_schemes))


async def _locate_quant(client, entry, scheme_rows):
    as_of = _locate_month_end()
    base = entry.endpoint_url.rsplit('/',1)[0] + '/'
    async def call(method, body):
        response = await client.post(base+method, content=body, headers={'Content-Type':'application/json; charset=utf-8'})
        return _response_json(response).get('d','')
    combined = await call('displaydisclouser', f"{{id:'{as_of.year}',cat:'MONTHLY PORTFOLIO'}}")
    records = _html_records(combined, entry.endpoint_url)
    months = await call('displaydisclouser1', f"{{id:'{as_of.year}',cat:'MONTHLY PORTFOLIO - FUND - WISE'}}")
    ids = re.findall(r'(?:\bid\s*=\s*["\']?|submit_event2\(\s*["\']?)(\d{1,2})\b', months)
    if str(as_of.month) in ids:
        per_fund = await call('displaydisclouser2', f"{{id:'{as_of.month}',cat:'MONTHLY PORTFOLIO - FUND - WISE',tab:'{as_of.year}'}}")
        records.extend(_html_records(per_fund, entry.endpoint_url, hints=True))
    return await _fetch_dated_records(client, records, as_of)


async def _locate_abakkus(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url); page.raise_for_status()
    script = re.search(r'<script\b(?=[^>]*\bid=["\']verticals-data["\'])[^>]*>(.*?)</script>', page.text, re.I|re.S)
    if not script:
        return []
    records = []
    for vertical in json.loads(script.group(1)):
        if vertical.get('title') != 'Monthly Portfolio Disclosures':
            continue
        for item in _walk_dicts(vertical):
            media = item.get('downloadMedia') or {}
            url = item.get('downloadUrl') or media.get('url')
            if url:
                records.append((_file_url(url, entry.landing_url), _disclosed_date(item.get('title')), None))
    return await _fetch_dated_records(client, records, as_of, keep_all=False)


async def _locate_zerodha(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url); page.raise_for_status()
    records = [(_file_url(url, str(page.url)), _filename_as_of_date(url), None)
               for url, label in _extract_links(page.text) if _candidate_allowed(entry, url, label)]
    files = await _fetch_dated_records(client, records, as_of)
    result = []
    for file in files:
        if file.data is None:
            result.append(file)
            continue
        text = ' '.join(str(c or '') for row in _workbook_title_rows(file.data, 4) for c in row)
        title = re.search(r'MONTHLY PORTFOLIO STATEMENT OF\s+(.+?)\s+FOR\s+', text, re.I)
        result.append(LocatedFile(file.data,file.source_url,file.as_of_date,title.group(1).strip() if title else None))
    return result


async def _locate_unifi(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url); page.raise_for_status()
    records = []
    for url, label in _extract_links(page.text):
        if not _candidate_allowed(entry,url,label):
            continue
        filename = urllib.parse.unquote(urllib.parse.urlsplit(url).path.rsplit('/',1)[-1])
        name = re.fullmatch(r'MP-(Unifi-.*)-\d{8}\.xlsx?', filename, re.I)
        if name:
            records.append((_file_url(url,str(page.url)),_filename_as_of_date(url),name.group(1).replace('-',' ')))
    return await _fetch_dated_records(client,records,as_of)


async def _locate_tata(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url); page.raise_for_status()
    records = []
    for row in _page_objects(page.text):
        url = row.get('field_media_document')
        if isinstance(url,str) and 'monthly portfolio' in urllib.parse.unquote(url).lower():
            records.append((_file_url(url,str(page.url)),_disclosed_date(row.get('field_document_title')),None))
    return await _fetch_dated_records(client,records,as_of,keep_all=False)


async def _locate_sundaram(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url); page.raise_for_status()
    script = re.search(r'<script[^>]+src=["\']([^"\']*Modules_Disclosure_Monthly_Fortnightly_Adhoc_Portfolios,App_Web_[^"\']+\.ashx)["\']',page.text,re.I)
    if not script:
        return []
    response = await client.post(urllib.parse.urljoin(str(page.url),script.group(1)),params={'_method':'GetCategory','_session':'no'},
                                 content='Catid=Monthly',headers={'Content-Type':'text/plain'})
    response.raise_for_status()
    return await _fetch_dated_records(client,_html_records(response.text,str(response.url)),as_of)


async def _locate_bajaj(client, entry, scheme_rows):
    as_of = _locate_month_end()
    page = await client.get(entry.landing_url); page.raise_for_status()
    nonce = re.search(r'bajajDownloads\s*=\s*\{[^}]*?["\']nonce["\']\s*:\s*["\']([0-9a-f]+)',page.text)
    if not nonce:
        return []
    fy = as_of.year if as_of.month>=4 else as_of.year-1
    year = f'{fy}-{(fy+1)%100:02d}'
    common = {'nonce':nonce.group(1),'section_id':'757'}
    async def call(fields):
        response = await client.post(entry.endpoint_url,data={**common,**fields})
        return _response_json(response).get('data',{})
    years = await call({'action':'bajaj_get_filter_options','filter_for':'years'})
    if not any(str(r.get('value'))==year for r in years.get('options',[])):
        return []
    months = await call({'action':'bajaj_get_filter_options','filter_for':'months','year':year})
    month = as_of.strftime('%B')
    if not any(month.lower() in {str(r.get('value','')).lower(),str(r.get('label','')).lower()} for r in months.get('options',[])):
        return []
    data = await call({'action':'bajaj_get_downloads','year':year,'month':month})
    markup = data.get('html','')
    titles = re.findall(r'<span[^>]*class=["\'][^"\']*bd-download-title[^"\']*["\'][^>]*>(.*?)</span>',markup,re.I|re.S)
    links = _extract_links(markup)
    records = [(_file_url(url,entry.endpoint_url),_disclosed_date(re.sub(r'<[^>]+>',' ',titles[i])) if i<len(titles) else _filename_as_of_date(url),None)
               for i,(url,label) in enumerate(links)]
    return await _fetch_dated_records(client,records,as_of,keep_all=False)


def _monthly_manifest_schemes(value, period=None, identity=None):
    if isinstance(value, list):
        for child in value:
            yield from _monthly_manifest_schemes(child, period, identity)
    elif isinstance(value, dict):
        identity = {**(identity or {}), **{k:value[k] for k in ('schemeName','folderName') if k in value}}
        declared = str(value.get('period') or value.get('category') or '').lower()
        if declared in {'monthly','halfyearly','fortnightly'}:
            period = declared
        if period == 'monthly' and 'financialYears' in value and len(identity)==2:
            yield {**identity,'financialYears':value['financialYears']}
        for key, child in value.items():
            token = re.sub(r'[^a-z]', '', key.lower())
            child_period = token if token in {'monthly','halfyearly','fortnightly'} else period
            yield from _monthly_manifest_schemes(child, child_period, identity)


async def _locate_alphagrep(client, entry, scheme_rows):
    as_of = _locate_month_end()
    payload = await _json_request(client,entry,as_of)
    records = []
    for scheme in _monthly_manifest_schemes(payload):
        for year in scheme['financialYears']:
            for doc in year.get('documents',[]):
                filename = doc.get('fileName','')
                disclosed = _disclosed_date(filename)
                if not disclosed:
                    continue
                ext = '' if re.search(r'\.xlsx?$',filename,re.I) else '.xls'
                path = '/assets/documents/'+scheme['folderName']+'/monthly/'+year['yearFolder']+'/'+filename+ext
                records.append((_file_url(path,entry.endpoint_url),disclosed,scheme['schemeName']))
    return await _fetch_dated_records(client,records,as_of)


async def _locate_wealth(client, entry, scheme_rows):
    as_of = _locate_month_end()
    records, page = [], 1
    while True:
        response = await client.get(entry.landing_url,params={'page':page} if page>1 else None)
        response.raise_for_status()
        page_count = 1
        for row in _page_objects(response.text):
            if 'pageCount' in row:
                page_count=max(page_count,int(row['pageCount']))
            if not all(k in row for k in ('uploadDate','name','attachment')):
                continue
            name=html.unescape(row['name']).strip()
            disclosed=_disclosed_date(name)
            if disclosed is None:
                try:
                    disclosed=date.fromisoformat(row['uploadDate'])
                except (ValueError,TypeError):
                    continue
            hint=re.sub(r'^Monthly\s*-\s*','',name,flags=re.I)
            hint=re.sub(rf'\s*-\s*{_MONTH_PATTERN}\s+\d{{1,2}},?\s+\d{{4}}\s*$','',hint,flags=re.I).strip()
            url=(row.get('attachment') or {}).get('url')
            if isinstance(url,str) and _candidate_allowed(entry,url):
                records.append((_file_url(url,entry.landing_url),disclosed,hint))
        if page>=page_count:
            break
        page+=1
    return await _fetch_dated_records(client,records,as_of)


async def _locate_360(client, entry, scheme_rows):
    as_of = _locate_month_end()
    candidates=await _static_regex_candidates(client,entry)
    records=[(_file_url(url,entry.landing_url),_filename_as_of_date(url),None) for url,label in candidates]
    return await _fetch_dated_records(client,records,as_of,keep_all=False)


async def _locate_motilal(client, entry, scheme_rows):
    as_of=_locate_month_end()
    sitemap=await client.get(entry.landing_url);sitemap.raise_for_status()
    pages=dict.fromkeys(re.findall(r'<loc>\s*(https?://[^<]+/mutual-funds/[^<]+)\s*</loc>',sitemap.text))
    records=[]
    for url in pages:
        response=await client.get(html.unescape(url));response.raise_for_status()
        text=html.unescape(re.sub(r'<[^>]+>',' ',response.text))
        path=re.search(r'portfolioUrl\s*:\s*(/\S+\.xlsx?)',text,re.I)
        dated=re.search(r'Date As On\s*:\s*(\d{1,2}\s+[A-Za-z]+\s+\d{4})',text,re.I)
        if path and dated:
            records.append((_file_url(path.group(1),str(response.url)),_disclosed_date(dated.group(1)),None))
    files=await _fetch_dated_records(client,records,as_of)
    result=[]
    for file in files:
        if file.data is None:
            result.append(file)
            continue
        rows=_workbook_title_rows(file.data)
        title=next((str(c).strip() for c in rows[7] if c and 'motilal' in str(c).lower()),None) if len(rows)>7 else None
        result.append(LocatedFile(file.data,file.source_url,file.as_of_date,title))
    return result


async def _locate_choice(client, entry, scheme_rows):
    as_of=_locate_month_end()
    payload=await _json_request(client,entry,as_of)
    records=[]
    for root in payload.get('body',[]):
        if root.get('slug')!='regulatory-compliance':
            continue
        for child in root.get('children',[]):
            if child.get('slug') not in {'monthly-portfolio','fortnight-portfolio'}:
                continue
            for node in _walk_dicts(child):
                url=node.get('open_file_url')
                if isinstance(url,str):
                    records.append((_file_url(url,entry.endpoint_url),_filename_as_of_date(url),node.get('scheme_name')))
    return await _fetch_dated_records(client,records,as_of)


async def _locate_quantum(client, entry, scheme_rows):
    as_of = _locate_month_end()
    candidates = await _static_regex_candidates(client, entry)
    records = [(_file_url(url,entry.landing_url),_disclosed_date(label),None)
               for url,label in candidates if re.search(r'\bAll Funds\b',label,re.I)]
    return await _fetch_dated_records(client,records,as_of,keep_all=False)


async def _locate_nj(client, entry, scheme_rows):
    as_of = _locate_month_end()
    records = []
    for url,label in await _static_regex_candidates(client,entry):
        filename = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query).get('file',[''])[0]
        records.append((_file_url(url,entry.landing_url),_filename_as_of_date('https://filename.invalid/'+urllib.parse.quote(filename)),None))
    files = await _fetch_dated_records(client,records,as_of)
    result = []
    for file in files:
        if file.data is None:
            result.append(file)
            continue
        rows = _workbook_title_rows(file.data,2)
        title = next((str(c).strip().splitlines()[0] for c in rows[1] if c and re.match(r'^NJ\s',str(c).strip())),None) if len(rows)>1 else None
        result.append(LocatedFile(file.data,file.source_url,file.as_of_date,title))
    return result


_CUSTOM_LOCATORS = {
    'franklin': _locate_franklin, 'axis': _locate_axis, 'lic': _locate_lic,
    'sbi': _locate_sbi, 'trust': _locate_trust, 'mirae': _locate_mirae,
    'boi': _locate_boi, 'taurus': _locate_taurus, 'absl': _locate_absl,
    'navi': _locate_navi, 'pgim': _locate_pgim, 'whiteoak': _locate_whiteoak,
    'icici': _locate_icici, 'uti': _locate_uti,
    'jio': _locate_jio, 'quant': _locate_quant, 'abakkus': _locate_abakkus,
    'zerodha': _locate_zerodha, 'unifi': _locate_unifi, 'tata': _locate_tata,
    'sundaram': _locate_sundaram, 'bajaj': _locate_bajaj, 'alphagrep': _locate_alphagrep,
    'wealth': _locate_wealth, '360': _locate_360, 'motilal': _locate_motilal, 'choice': _locate_choice,
    'quantum': _locate_quantum, 'nj': _locate_nj,
}


async def _cdx_query(client: httpx.AsyncClient, url: str, *, match_type: str = "exact") -> list[str]:
    """Return captured originals, not archive download URLs.

    limit=-50 selects the last 50 index rows, not the latest 50 disclosures;
    older months and some schemes may be absent. Dates are ranked separately.
    """
    resp = await client.get("https://web.archive.org/cdx/search/cdx", params={
        "url": url, "output": "json", "matchType": match_type,
        "filter": "statuscode:200", "collapse": "urlkey", "limit": "-50",
    })
    resp.raise_for_status()
    rows = resp.json()
    if len(rows) <= 1:
        return []
    header, *data_rows = rows
    original_idx = header.index("original")
    return [row[original_idx] for row in data_rows]


async def _locate_wayback_cached_page(client: httpx.AsyncClient, entry: DisclosureResolverEntry) -> list[LocatedFile]:
    if not await _cdx_query(client, entry.landing_url, match_type="exact"):
        return []
    cached = await client.get(f"https://web.archive.org/web/2id_/{entry.landing_url}")
    cached.raise_for_status()
    candidates = []
    for href, label in _extract_links(cached.text):
        if not _candidate_allowed(entry, href, label):
            continue
        absolute = urllib.parse.urljoin(entry.landing_url, href)
        parsed = urllib.parse.urlsplit(absolute)
        live_url = parsed._replace(netloc=entry.cdn_domain, scheme="https").geturl()
        candidates.append((live_url, label))
    return await _fetch_candidates(client, candidates, keep_all=True)


async def _locate_wayback_prefix_search(client: httpx.AsyncClient, entry: DisclosureResolverEntry) -> list[LocatedFile]:
    originals = await _cdx_query(client, entry.landing_url, match_type="prefix")
    # CDX rows include PDFs, images and pages under the prefix: only workbook/zip paths are candidates.
    candidates = [(url, url.rsplit("/", 1)[-1]) for url in originals
                  if _WORKBOOK_URL.search(urllib.parse.unquote(url)) and _candidate_allowed(entry, url)]
    return await _fetch_candidates(client, candidates, keep_all=True)


async def _locate_wayback(client: httpx.AsyncClient, entry: DisclosureResolverEntry) -> list[LocatedFile]:
    if entry.cdn_domain:
        return await _locate_wayback_cached_page(client, entry)
    return await _locate_wayback_prefix_search(client, entry)


def _alert(amc_name: str, reason: str, detail: str) -> None:
    logger.warning("PORTFOLIO_DISCLOSURE_ALERT amc=%s reason=%s detail=%s",
                   " ".join(amc_name.split()), " ".join(reason.split()), " ".join(detail.split()))


def _already_written(db: Session, amc_name: str, as_of_date: date) -> bool:
    return db.query(DisclosureBatch).filter_by(
        amc_name=amc_name, as_of_date=as_of_date, status=DisclosureBatchStatus.WRITTEN).first() is not None


def _written_source_urls(db: Session, amc_name: str, as_of: date) -> set[str]:
    return {url for (url,) in db.query(DisclosureBatch.source_url).filter_by(
        amc_name=amc_name, as_of_date=as_of, status=DisclosureBatchStatus.WRITTEN) if url}


def _alert_due(today: date) -> bool:
    return today.day in _ALERT_DAYS


def monthly_zero_batch_summary(db: Session, today: date) -> list[str]:
    represented = {amc for (amc,) in db.query(Scheme.amc_name).distinct()}
    written = {amc for (amc,) in db.query(DisclosureBatch.amc_name).filter_by(
        as_of_date=_latest_month_end(today), status=DisclosureBatchStatus.WRITTEN).distinct()}
    return sorted(set(AMC_DISCLOSURE_RESOLVERS).intersection(represented) - written)


def _downloaded_hint(data: bytes, entry: DisclosureResolverEntry, hint: str | None) -> str | None:
    # Exact title readers moved mechanically from discovery's three workbook readers.
    if entry.custom_locator == 'zerodha':
        text = ' '.join(str(c or '') for row in _workbook_title_rows(data, 4) for c in row)
        title = re.search(r'MONTHLY PORTFOLIO STATEMENT OF\s+(.+?)\s+FOR\s+', text, re.I)
        return title.group(1).strip() if title else None
    if entry.custom_locator == 'motilal':
        rows = _workbook_title_rows(data)
        return next((str(c).strip() for c in rows[7] if c and 'motilal' in str(c).lower()), None) if len(rows) > 7 else None
    if entry.custom_locator == 'nj':
        rows = _workbook_title_rows(data, 2)
        return next((str(c).strip().splitlines()[0] for c in rows[1] if c and re.match(r'^NJ\s', str(c).strip())), None) if len(rows) > 1 else None
    return hint


async def _iter_located_files(client, entry, scheme_rows, skip_urls: set[str]) -> AsyncIterator[LocatedFile]:
    """Discover metadata first, then hold one downloaded body until the caller finishes.

    Failure references (data=None) let the caller record failures without hiding peers.
    No WRITTEN URL is requested. Discovery context is task-local and always reset.
    """
    token = _DISCOVERY_ONLY.set(True)
    try:
        references = await _locate_and_fetch(client, entry, scheme_rows)
    finally:
        _DISCOVERY_ONLY.reset(token)
    single = entry.file_shape is DisclosureFileShape.SINGLE_WORKBOOK_MULTI_SHEET
    seen = set()
    for reference in references:
        if reference.source_url in seen:
            continue
        seen.add(reference.source_url)
        if reference.source_url in skip_urls:
            if single:
                break
            continue
        error = None
        try:
            data = await _download_workbook(client, reference.source_url)
        except httpx.HTTPError as exc:
            error = f'{type(exc).__name__}: {exc}'
            data = None
        if data is None:
            _alert(scheme_rows[0].amc_name, 'error', f'file={reference.source_url} {error or "unavailable or invalid workbook"}')
            yield reference
            continue
        try:
            hint = _downloaded_hint(data, entry, reference.scheme_hint)
        except Exception as exc:
            # Retain the bytes for Parse/archive and continue to later files.
            _alert(scheme_rows[0].amc_name, 'parse_error', f'file={reference.source_url} title: {exc}')
            hint = reference.scheme_hint
        yield LocatedFile(data, reference.source_url, reference.as_of_date, hint)
        del data
        if single:
            break


def _zip_members(data: bytes, amc_name: str):
    # Never materialize every decompressed member (ICICI has 151 workbooks).
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        for name in archive.namelist():
            if not name.lower().endswith(('.xls', '.xlsx')) or name.endswith('/'):
                continue
            try:
                member_data = archive.read(name)
            except Exception as exc:
                _alert(amc_name, 'parse_error', f'member={name} {type(exc).__name__}: {exc}')
                yield name, None
                continue
            yield name, member_data
            del member_data


def _parse_file(file: LocatedFile, amc_name: str):
    """Yield member name and parsed sheets; an invalid ZIP member cannot stop peers."""
    try:
        sniff_workbook_format(file.data)
    except ValueError:
        if not zipfile.is_zipfile(io.BytesIO(file.data)):
            raise
        parsed_member = False
        for name, data in _zip_members(file.data, amc_name):
            if data is None:
                yield name, [], {'__parse_error__': 'unreadable ZIP member'}
                continue
            try:
                sheets = parse_workbook(data)
                index_map = read_index_sheet(data)
            except Exception as exc:
                _alert(amc_name, "parse_error", f"member={name} {type(exc).__name__}: {exc}")
                yield name, [], {'__parse_error__': str(exc)}
                continue
            parsed_member = True
            yield name, sheets, index_map
            del data, sheets, index_map
        if not parsed_member:
            raise ValueError("ZIP contains no readable workbook members")
    else:
        # A broken standalone workbook is an AMC error, not an isolated ZIP peer.
        yield file.source_url, parse_workbook(file.data), read_index_sheet(file.data)


async def refresh_portfolio_disclosures(db: Session, today: date | None = None) -> PortfolioDisclosureRefreshResult:
    started = time.perf_counter()
    effective_today = today or date.today()
    target = _latest_month_end(today) if today is not None else _latest_month_end()
    token = _RUN_TODAY.set(today)
    amcs_processed = amcs_failed = amcs_not_onboarded = schemes_written = 0
    try:
        async with httpx.AsyncClient(timeout=_HTTP_TIMEOUT, headers=_HEADERS, follow_redirects=True) as http_client:
            for amc_name, entry in AMC_DISCLOSURE_RESOLVERS.items():
                if entry.kind is DisclosureResolverKind.MANUAL:
                    if _alert_due(effective_today):
                        waiting = db.query(DisclosureBatch).filter_by(amc_name=amc_name, as_of_date=target,
                            status=DisclosureBatchStatus.AWAITING_MANUAL_UPLOAD).count()
                        if waiting:
                            _alert(amc_name, 'manual_upload_pending', f'waiting={waiting}')
                    continue
                if not _is_locate_implemented(entry):
                    amcs_not_onboarded += 1
                    continue
                try:
                    scheme_rows = db.query(Scheme).filter(Scheme.amc_name == amc_name).all()
                    if not scheme_rows:
                        _alert(amc_name, 'zero_schemes_matched', 'registry key has no scheme_master rows')
                        amcs_failed += 1
                        continue
                    families = build_families(scheme_rows)
                    skip_urls = _written_source_urls(db, amc_name, target)
                    files_seen = written_this_amc = 0
                    failed = False
                    newest = None
                    failed_newest = None
                    unmatched_names = []
                    async for file in _iter_located_files(http_client, entry, scheme_rows, skip_urls):
                        files_seen += 1
                        newest = max(newest or file.as_of_date, file.as_of_date)
                        if file.data is None:
                            failed = True
                            failed_newest = max(failed_newest or file.as_of_date, file.as_of_date)
                            del file
                            continue
                        if (entry.file_shape is DisclosureFileShape.SINGLE_WORKBOOK_MULTI_SHEET
                                and failed_newest and file.as_of_date < failed_newest):
                            _alert(amc_name, 'stale_fallback', f'newest={failed_newest} used={file.as_of_date}')
                        file_failed = False
                        failed_batches = []
                        written_this_file = 0
                        hinted_sheet_used = False
                        try:
                            with db.begin_nested():
                                for member, sheets, index_map in _parse_file(file, amc_name):
                                    if '__parse_error__' in index_map:
                                        file_failed = True
                                    for sheet in sheets:
                                        if sheet.error:
                                            file_failed = True
                                            _alert(amc_name, 'parse_error', f'file={member} sheet={sheet.sheet_name} {sheet.error}')
                                        if not sheet.holdings:
                                            continue
                                        if file.scheme_hint:
                                            if hinted_sheet_used:
                                                continue
                                            hinted_sheet_used = True
                                        parse_errors = ([{'file': member, 'sheet': sheet.sheet_name, 'error': sheet.error}]
                                                        if sheet.error else None)
                                        match_key = sheet.scheme_title or file.scheme_hint or index_map.get(sheet.sheet_name, sheet.sheet_name)
                                        match = match_scheme_name(match_key, families)
                                        if match is None:
                                            if match_key not in unmatched_names and len(unmatched_names) < 5:
                                                unmatched_names.append(match_key)
                                            continue
                                        family, method, confidence = match
                                        written_this_file += write_disclosure_batch(db, family, file.as_of_date,
                                            amc_name, file.source_url, sheet.holdings, method, confidence,
                                            parse_errors=parse_errors, declared_total=sheet.declared_total)
                                        batches = db.query(DisclosureBatch).filter(
                                            DisclosureBatch.scheme_id.in_([scheme.id for scheme in family.schemes]),
                                            DisclosureBatch.as_of_date == file.as_of_date).all()
                                        for batch in batches:
                                            if batch.status is DisclosureBatchStatus.FAILED:
                                                file_failed = True
                                                failed_batches.append(batch)
                                        if any(batch.status is DisclosureBatchStatus.FAILED for batch in batches):
                                            _alert(amc_name, 'failed_sum_check', f'file={file.source_url} sheet={sheet.sheet_name}')
                        except Exception as exc:
                            file_failed = True
                            written_this_file = 0
                            failed_batches = []
                            failed = True
                            _alert(amc_name, 'error', f'file={file.source_url} {type(exc).__name__}: {exc}'[:300])
                        if file_failed:
                            location = archive_failed_file(file.data, amc_name, file.as_of_date,
                                urllib.parse.unquote(urllib.parse.urlsplit(file.source_url).path.rsplit('/', 1)[-1]))
                            for batch in failed_batches:
                                batch.raw_file_reference = location
                        written_this_amc += written_this_file
                        del file
                    if unmatched_names:
                        _alert(amc_name, 'unmatched_scheme', ' | '.join(unmatched_names))
                    if not files_seen and not skip_urls:
                        failed = True
                        if _alert_due(effective_today):
                            _alert(amc_name, 'no_file_found', entry.landing_url or entry.url_template or '')
                    if (entry.kind is DisclosureResolverKind.WAYBACK_LOCATE and newest
                            and newest < target and _alert_due(effective_today)):
                        _alert(amc_name, 'stale_wayback_capture', f'newest={newest}')
                    await commit_off_loop(db)
                    schemes_written += written_this_amc
                    if failed:
                        amcs_failed += 1
                    else:
                        amcs_processed += 1
                except Exception as exc:
                    db.rollback()
                    _alert(amc_name, 'error', f'{type(exc).__name__}: {exc}'[:300])
                    amcs_failed += 1
    finally:
        _RUN_TODAY.reset(token)
    return PortfolioDisclosureRefreshResult(
        success=True, amcs_processed=amcs_processed, amcs_failed=amcs_failed,
        amcs_not_onboarded=amcs_not_onboarded, schemes_written=schemes_written,
        seconds=round(time.perf_counter() - started, 1),
    )
