import asyncio
import json
import io
import logging
import uuid
import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import httpx
import openpyxl
import pytest

from app.models.enums import DisclosureBatchStatus
from app.models.reference import DisclosureBatch, Scheme, SchemeHolding
from app.services.analytics import portfolio_disclosure_client as client
from app.services.analytics.portfolio_disclosure_resolvers import (
    AMC_DISCLOSURE_RESOLVERS, DisclosureFileShape as Shape,
    DisclosureResolverEntry as Entry, DisclosureResolverKind as Kind,
)


@pytest.fixture
def mock_http(monkeypatch):
    """Every refresh request must go through MockTransport; unexpected URLs fail."""
    real_client = httpx.AsyncClient
    requests = []

    def install(routes):
        def handler(request):
            requests.append(str(request.url))
            key = str(request.url)
            assert key in routes, f"Unexpected HTTP request: {key}"
            response = routes[key]
            return response(request) if callable(response) else response

        monkeypatch.setattr(client.httpx, "AsyncClient", lambda **kw: real_client(
            transport=httpx.MockTransport(handler), trust_env=False, **kw))
        return requests

    return install


def workbook(sheets=None):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, weight, total in sheets or [("Example Equity Fund", 100, 100)]:
        ws = wb.create_sheet(name)
        ws.append(["Name of the Instrument", "ISIN", "Industry", "Quantity",
                   "Market Value (Rs. in Lakhs)", "% to NAV"])
        ws.append(["Equity & Equity related"])
        ws.append(["HDFC Bank Ltd", "INE040A01034", "Banks", 1000, 15, weight])
        if total is not None:
            ws.append(["Grand Total", None, None, None, 15, total])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def seed(db, amc="Example Mutual Fund", base="Example Equity Fund", plans=1):
    rows = [Scheme(id=uuid.uuid4(), name=f"{base} - Plan {i}", amc_name=amc,
                   base_name=base, sebi_category="Equity Scheme - Flexi Cap Fund")
            for i in range(plans)]
    db.add_all(rows)
    db.commit()
    return rows


def static_entry(shape=Shape.SINGLE_WORKBOOK_MULTI_SHEET):
    return Entry(kind=Kind.STATIC_REGEX, file_shape=shape,
                 landing_url="https://example.test/list", link_pattern="portfolio",
                 exclude_pattern="fortnightly")


def run_refresh(db, monkeypatch, mock_http, routes, entry=None, amc="Example Mutual Fund"):
    monkeypatch.setattr(client, "AMC_DISCLOSURE_RESOLVERS", {amc: entry or static_entry()})
    requests = mock_http(routes)
    return asyncio.run(client.refresh_portfolio_disclosures(db)), requests


@pytest.mark.parametrize("filename,expected", [
    ("NIMF-MONTHLY-PORTFOLIO-31-Aug-26.xls", date(2026, 8, 31)),
    ("helios-flexi-cap-fund-monthly-portfolio-as-on-30th-september-2026.xls", date(2026, 9, 30)),
    ("helios-overnight-fund-monthly-portfolio-as-on-30st-september-2026.xls", date(2026, 9, 30)),
    ("MP-Unifi-Flexi-Cap-Fund-30092026.xlsx", date(2026, 9, 30)),
    ("MP-Unifi-Liquid-Fund-31082026.xls", date(2026, 8, 31)),
    ("OBMF_Flexi_Cap_Portfolio_Sept_26_0cfdd46aac.xlsx", date(2026, 9, 30)),
    ("Old_Bridge_Arbitrage_Fund_January_2026_deadbeef.xlsx", date(2026, 1, 31)),
    ("IN_MF_MONTHLY_PORTFOLIO_Sep2026_Final_fc91906e39.xls", date(2026, 9, 30)),
    ("IN_MF_MONTHLY_PORTFOLIO_Apr_2026_final_x.xlsx", date(2026, 4, 30)),
    ("Monthly%20HDFC%20Flexi%20Cap%20Fund%20-%2031%20August%202026.xlsx", date(2026, 8, 31)),
    ("SC-%E2%80%93-Canara-Robeco-Small-Cap-Fund-%E2%80%93-September-2026.xlsx", date(2026, 9, 30)),
    ("OF-Canara-Robeco-Overnight-Fund-September-2026.xlsx", date(2026, 9, 30)),
    ("portfolio-30-Sep-26.xlsx", date(2026, 9, 30)),
    ("portfolio-30-09-2026.xlsx", date(2026, 9, 30)),
    ("portfolio-February-2024.xlsx", date(2024, 2, 29)),
    ("portfolio-31-Feb-26.xls", None),
    ("portfolio-31022026.xls", None),
    ("portfolio-30-Foo-2026.xlsx", None),
    ("current.xls", None),
    ("Monthly_Portfolio_Flexi_Cap_0fd0dc18c8.xlsx", None),
])
def test_filename_as_of_date_real_shapes_and_invalid_dates(filename, expected):
    # Dates must come from the filename, never the upload folder or query string.
    assert client._filename_as_of_date(f"https://example.test/2026/10/{filename}?date=30-Sep-26") == expected


@pytest.mark.parametrize("snippet,url", [
    (r'<script>{\"url\":\"https://files.hdfcfund.com/s3fs-public/2026-09/Monthly%20HDFC%20Flexi%20Cap%20Fund%20-%2031%20August%202026.xlsx\"}</script>',
     "https://files.hdfcfund.com/s3fs-public/2026-09/Monthly%20HDFC%20Flexi%20Cap%20Fund%20-%2031%20August%202026.xlsx"),
    (r'<script>{\"attachment\":{\"url\":\"/uploads/Monthly_Portfolio_Flexi_Cap_0fd0dc18c8.xlsx\"}}</script>',
     "/uploads/Monthly_Portfolio_Flexi_Cap_0fd0dc18c8.xlsx"),
    (r'<script>{\"fileUrl\":\"https://s3.ap-south-1.amazonaws.com/x-web-s3.360.one/IN_MF_MONTHLY_PORTFOLIO_Sep2026_Final_fc91906e39.xls\"}</script>',
     "https://s3.ap-south-1.amazonaws.com/x-web-s3.360.one/IN_MF_MONTHLY_PORTFOLIO_Sep2026_Final_fc91906e39.xls"),
    (r'<script>{"url":"https:\/\/example.test\/portfolio-30-Sep-26.xlsx?a=1\u0026b=2"}</script>',
     "https://example.test/portfolio-30-Sep-26.xlsx?a=1&b=2"),
])
def test_extract_links_unescapes_trimmed_capture_payloads(snippet, url):
    assert client._extract_links(snippet) == [(url, "")]


def test_extract_links_preserves_anchor_label_and_deduplicates():
    snippet = '<a href="/portfolio-30-Sep-26.xls?a=1&amp;b=2"><b>Example</b> Equity Fund</a>'
    assert client._extract_links(snippet) == [("/portfolio-30-Sep-26.xls?a=1&b=2", "Example Equity Fund")]


def test_refresh_latest_file_magic_bytes_plan_variants_and_rerun(db_session, monkeypatch, mock_http):
    seed(db_session, plans=2)
    old = "https://example.test/portfolio-31-Aug-26.xls"
    new = "https://example.test/portfolio-30-Sep-26.xls"  # OOXML under .xls
    routes = {
        "https://example.test/list": httpx.Response(200, text=f'<a href="{old}">Old</a><a href="{new}">New</a>'),
        new: httpx.Response(200, content=workbook()),
    }
    result, requests = run_refresh(db_session, monkeypatch, mock_http, routes)
    assert (result.success, result.amcs_processed, result.amcs_failed, result.schemes_written) == (True, 1, 0, 2)
    assert old not in requests
    batches = db_session.query(DisclosureBatch).all()
    assert len(batches) == 2 and all(b.status is DisclosureBatchStatus.WRITTEN for b in batches)
    assert all(b.as_of_date == date(2026, 9, 30) for b in batches)
    assert db_session.query(SchemeHolding).first().market_value == Decimal("1500000")
    result2 = asyncio.run(client.refresh_portfolio_disclosures(db_session))
    assert result2.amcs_processed == 1 and result2.schemes_written == 0
    assert db_session.query(SchemeHolding).count() == 2


def test_multifile_matches_link_text_and_excludes_fortnightly(db_session, monkeypatch, mock_http):
    seed(db_session, base="Helios Flexi Cap Fund")
    seed(db_session, base="Helios Small Cap Fund")
    urls = [f"https://example.test/helios-{name}-fund-monthly-portfolio-as-on-30st-september-2026.xls"
            for name in ("flexi-cap", "small-cap")]
    html = ''.join(f'<a href="{u}">{label} - Monthly Portfolio</a>' for u, label in zip(
        urls, ("Helios Flexi Cap Fund", "Helios Small Cap Fund")))
    html += '<a href="/fortnightly-portfolio-30-Sep-26.xls">Fortnightly</a>'
    routes = {"https://example.test/list": httpx.Response(200, text=html)}
    routes.update({u: httpx.Response(200, content=workbook([("Sheet", 100, 100)])) for u in urls})
    result, requests = run_refresh(db_session, monkeypatch, mock_http, routes,
                                   static_entry(Shape.MULTI_FILE_SINGLE_SCHEME))
    assert result.schemes_written == 2
    assert len(requests) == 3
    assert db_session.query(DisclosureBatch).count() == 2


def test_template_uses_month_end_and_deduplicates_plan_variants(db_session, monkeypatch, mock_http):
    seed(db_session, base="Union Corporate Bond Fund", plans=2)
    seed(db_session, base="Union Missing Fund")
    monkeypatch.setattr(client, "_latest_month_end", lambda: date(2026, 9, 30))
    entry = Entry(kind=Kind.STATIC_REGEX, file_shape=Shape.DETERMINISTIC_TEMPLATE,
                  url_template="https://example.test/{month_name_lower}-{yyyy}/{scheme_slug}-{dd}-{mm}-{yyyy}.xlsx")
    good = "https://example.test/september-2026/union-corporate-bond-fund-30-09-2026.xlsx"
    bad = "https://example.test/september-2026/union-missing-fund-30-09-2026.xlsx"
    result, requests = run_refresh(db_session, monkeypatch, mock_http, {
        good: httpx.Response(200, content=workbook([("Sheet", 100, 100)])), bad: httpx.Response(500),
    }, entry)
    assert result.schemes_written == 2
    assert requests == [good, bad]


def test_template_supplies_all_fields_and_fetches_amc_bundle_once(db_session, monkeypatch, mock_http):
    seed(db_session, base="Example Equity Fund", plans=2)
    seed(db_session, base="Example Debt Fund")
    monkeypatch.setattr(client, "_latest_month_end", lambda: date(2026, 9, 30))
    entry = Entry(kind=Kind.STATIC_REGEX, file_shape=Shape.DETERMINISTIC_TEMPLATE,
                  url_template="https://example.test/{month_name}-{yyyy}-{yy}-{yy_plus_1}-{dd}-{mm}.xls")
    url = "https://example.test/September-2026-26-27-30-09.xls"
    result, requests = run_refresh(db_session, monkeypatch, mock_http, {
        url: httpx.Response(200, content=workbook([
            ("Example Equity Fund", 100, 100), ("Example Debt Fund", 100, 100)])),
    }, entry)
    assert requests == [url]
    assert result.schemes_written == 3


@pytest.mark.parametrize("weight,total,status", [
    (98, 100, DisclosureBatchStatus.FAILED),  # in band but unreconciled
    (100, 100, DisclosureBatchStatus.WRITTEN),
    (1, "100%", DisclosureBatchStatus.WRITTEN),
    (100, 1, DisclosureBatchStatus.FAILED),  # error WITH holdings
])
def test_refresh_passes_declared_total_and_keeps_errored_holdings(db_session, monkeypatch, mock_http, weight, total, status):
    seed(db_session)
    url = "https://example.test/portfolio-30-Sep-26.xls"
    result, _ = run_refresh(db_session, monkeypatch, mock_http, {
        "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
        url: httpx.Response(200, content=workbook([("Example Equity Fund", weight, total)])),
    })
    batch = db_session.query(DisclosureBatch).one()
    assert batch.status is status
    assert db_session.query(SchemeHolding).count() == 1
    if weight == 100 and total == 1:
        assert "scale disagrees" in str(batch.parse_errors)
    assert result.schemes_written == 1  # writer counts persisted scheme batches, including FAILED


def test_zip_bad_member_isolated_and_format_sniffed_despite_extension(db_session, monkeypatch, mock_http, caplog):
    seed(db_session)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("bad.xlsx", b"not a workbook")
        archive.writestr("valid.xls", workbook())
        archive.writestr("readme.txt", b"not holdings")
    url = "https://example.test/portfolio-30-Sep-26.xls"  # ZIP under .xls
    with caplog.at_level(logging.WARNING):
        result, _ = run_refresh(db_session, monkeypatch, mock_http, {
            "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
            url: httpx.Response(200, content=buf.getvalue()),
        }, static_entry(Shape.ZIP_OF_WORKBOOKS))
    assert result.amcs_processed == 1 and result.schemes_written == 1
    assert db_session.query(DisclosureBatch).one().status is DisclosureBatchStatus.WRITTEN
    assert "bad.xlsx" in caplog.text


def test_sheet_error_without_holdings_does_not_fail_good_sheets(db_session, monkeypatch, mock_http):
    seed(db_session)
    data = workbook([("Broken Empty", 100, 100), ("Example Equity Fund", 100, 100)])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb["Broken Empty"].cell(1, 5, "Market Value (Rs. in Millions)")
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    url = "https://example.test/portfolio-30-Sep-26.xlsx"
    result, _ = run_refresh(db_session, monkeypatch, mock_http, {
        "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
        url: httpx.Response(200, content=buf.getvalue()),
    })
    assert result.schemes_written == 1
    batch = db_session.query(DisclosureBatch).one()
    # Per-sheet errors: the empty broken sheet is not written and must not FAIL its good neighbour.
    assert batch.status is DisclosureBatchStatus.WRITTEN
    assert batch.parse_errors is None


def test_zero_schemes_matched_has_distinct_alert_and_no_http(db_session, monkeypatch, mock_http, caplog):
    with caplog.at_level(logging.WARNING):
        result, requests = run_refresh(db_session, monkeypatch, mock_http, {})
    assert result.amcs_failed == 1 and requests == []
    assert "reason=zero_schemes_matched" in caplog.text
    assert "no_file_found" not in caplog.text


def test_manual_and_unbuilt_json_are_skips_without_alerts_or_http(db_session, monkeypatch, mock_http, caplog):
    monkeypatch.setattr(client, "AMC_DISCLOSURE_RESOLVERS", {
        "Manual Mutual Fund": Entry(kind=Kind.MANUAL),
        "JSON Mutual Fund": Entry(kind=Kind.JSON_API, endpoint_url="https://example.test/api"),
    })
    requests = mock_http({})
    with caplog.at_level(logging.WARNING):
        result = asyncio.run(client.refresh_portfolio_disclosures(db_session))
    assert result.amcs_not_onboarded == 1 and result.amcs_failed == 0
    assert requests == [] and "PORTFOLIO_DISCLOSURE_ALERT" not in caplog.text


def test_no_dated_file_raises_no_file_found_without_guessing(db_session, monkeypatch, mock_http, caplog):
    seed(db_session)
    with caplog.at_level(logging.WARNING):
        result, requests = run_refresh(db_session, monkeypatch, mock_http, {
            "https://example.test/list": httpx.Response(200, text='<a href="/portfolio-current.xls">September 2026</a>'),
        })
    assert result.amcs_failed == 1 and len(requests) == 1
    assert "reason=no_file_found" in caplog.text


def test_fetch_candidates_newest_first_with_http_failure_fallback():
    requests = []
    def handler(request):
        requests.append(request.url.path)
        return httpx.Response(404) if "Sep" in request.url.path else httpx.Response(200, content=workbook())
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._fetch_candidates(http, [
                ("https://example.test/portfolio-31-Aug-26.xls", "August"),
                ("https://example.test/portfolio-30-Sep-26.xls", "September"),
            ], keep_all=False)
    files = asyncio.run(run())
    assert requests == ["/portfolio-30-Sep-26.xls", "/portfolio-31-Aug-26.xls"]
    assert files[0].as_of_date == date(2026, 8, 31)


def test_download_rejects_oversized_stream_and_accepts_boundary(monkeypatch):
    monkeypatch.setattr(client, "_MAX_WORKBOOK_BYTES", 5)
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(
                200, content=b"PK\x03\x0412" if r.url.path == "/big" else b"PK\x03\x041"))) as http:
            return (await client._download_workbook(http, "https://example.test/big"),
                    await client._download_workbook(http, "https://example.test/exact"))
    assert asyncio.run(run()) == (None, b"PK\x03\x041")


_CDX_HEADER = ["urlkey", "timestamp", "original", "mimetype", "statuscode", "digest", "length"]


def cdx_rows(urls):
    return [_CDX_HEADER] + [["key", "20260910104807", u, "text/html", "200", "digest", "1"] for u in urls]


@pytest.mark.parametrize("rows,expected", [
    ([], []), ([_CDX_HEADER], []),
    (cdx_rows(["https://example.test/a", "https://example.test/b"]),
     ["https://example.test/a", "https://example.test/b"]),
])
def test_cdx_query_parameters_and_original_column(rows, expected):
    def handler(request):
        assert request.url.host == "web.archive.org"
        assert dict(request.url.params) == {
            "url": "https://example.test/uploads/", "output": "json", "matchType": "prefix",
            "filter": "statuscode:200", "collapse": "urlkey", "limit": "-50",
        }
        return httpx.Response(200, json=rows)
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._cdx_query(http, "https://example.test/uploads/", match_type="prefix")
    assert asyncio.run(run()) == expected


def test_wayback_hdfc_capture_log_url_and_embedded_august_link():
    entry = AMC_DISCLOSURE_RESOLVERS["HDFC Mutual Fund"]
    assert entry.landing_url == "https://www.hdfcfund.com/statutory-disclosure/portfolio/monthly-portfolio"
    live = "https://files.hdfcfund.com/s3fs-public/2026-09/Monthly%20HDFC%20Flexi%20Cap%20Fund%20-%2031%20August%202026.xlsx"
    cached = r'<script>{\"url\":\"https:\/\/files.hdfcfund.com\/s3fs-public\/2026-09\/Monthly%20HDFC%20Flexi%20Cap%20Fund%20-%2031%20August%202026.xlsx\"}</script>'
    requests = []
    def handler(request):
        requests.append(str(request.url))
        if request.url.path == "/cdx/search/cdx":
            assert request.url.params["url"] == entry.landing_url
            assert request.url.params["matchType"] == "exact"
            return httpx.Response(200, json=cdx_rows([entry.landing_url]))
        if str(request.url) == f"https://web.archive.org/web/2id_/{entry.landing_url}":
            return httpx.Response(200, text=cached)
        assert str(request.url) == live
        return httpx.Response(200, content=workbook([("HDFCEQ", 100, 100)]))
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_wayback(http, entry)
    files = asyncio.run(run())
    assert len(files) == 1 and files[0].source_url == live
    assert files[0].as_of_date == date(2026, 8, 31)
    assert len(requests) == 3  # no invented September direct-template fetch


def test_wayback_cached_page_rewrites_relative_path_to_live_cdn():
    entry = Entry(kind=Kind.WAYBACK_LOCATE, landing_url="https://blocked.test/list", cdn_domain="files.test")
    requests = []
    def handler(request):
        requests.append(str(request.url))
        if request.url.path == "/cdx/search/cdx":
            return httpx.Response(200, json=cdx_rows([entry.landing_url]))
        if request.url.host == "web.archive.org":
            return httpx.Response(200, text='<a href="/portfolio-30-Sep-26.xlsx">Example Equity Fund</a>')
        assert str(request.url) == "https://files.test/portfolio-30-Sep-26.xlsx"
        return httpx.Response(200, content=workbook())
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_wayback(http, entry)
    files = asyncio.run(run())
    assert files[0].source_url == "https://files.test/portfolio-30-Sep-26.xlsx"
    assert files[0].scheme_hint == "Example Equity Fund"


def test_wayback_no_capture_does_not_fetch_cached_or_live_page():
    requests = []
    def handler(request):
        requests.append(str(request.url))
        assert request.url.path == "/cdx/search/cdx"
        return httpx.Response(200, json=[])
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_wayback(http, AMC_DISCLOSURE_RESOLVERS["HDFC Mutual Fund"])
    assert asyncio.run(run()) == []
    assert len(requests) == 1


def test_wayback_canara_filters_unrelated_fortnightly_and_older_months():
    entry = AMC_DISCLOSURE_RESOLVERS["Canara Robeco Mutual Fund"]
    prefix = "https://www.canararobeco.com/wp-content/uploads/"
    new = prefix + "2026/10/SC-%E2%80%93-Canara-Robeco-Small-Cap-Fund-%E2%80%93-September-2026.xlsx"
    overnight = prefix + "2026/10/OF-Canara-Robeco-Overnight-Fund-September-2026.xlsx"
    old = prefix + "2026/09/SC-Canara-Robeco-Small-Cap-Fund-August-2026.xlsx"
    fortnightly = prefix + "2026/09/BP-%E2%80%93-Canara-Robeco-Banking-and-PSU-Debt-Fund-%E2%80%93-15th-September-2026.xlsx"
    unrelated = prefix + "2026/10/logo.png"
    requests = []
    def handler(request):
        requests.append(str(request.url))
        if request.url.path == "/cdx/search/cdx":
            assert request.url.params["matchType"] == "prefix"
            # limit=-50 can omit older months: discovery makes no completeness claim.
            assert request.url.params["limit"] == "-50"
            return httpx.Response(200, json=cdx_rows([old, unrelated, fortnightly, new, overnight]))
        assert str(request.url) in {new, overnight}
        return httpx.Response(200, content=workbook())
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_wayback(http, entry)
    files = asyncio.run(run())
    assert {f.source_url for f in files} == {new, overnight}
    assert all(f.as_of_date == date(2026, 9, 30) for f in files)
    assert len(requests) == 3


def test_wayback_prefix_without_pattern_preserves_candidate_behavior():
    entry = Entry(kind=Kind.WAYBACK_LOCATE, landing_url="https://example.test/uploads/")
    url = "https://example.test/uploads/portfolio-30-Sep-26.xls"
    def handler(request):
        if request.url.host == "web.archive.org":
            return httpx.Response(200, json=cdx_rows([url]))
        assert str(request.url) == url
        return httpx.Response(200, content=workbook())
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_wayback(http, entry)
    assert asyncio.run(run())[0].source_url == url


def test_one_amc_http_failure_does_not_stop_next(db_session, monkeypatch, mock_http, caplog):
    seed(db_session, amc="Bad Mutual Fund", base="Bad Equity Fund")
    seed(db_session)
    monkeypatch.setattr(client, "AMC_DISCLOSURE_RESOLVERS", {
        "Bad Mutual Fund": Entry(kind=Kind.STATIC_REGEX, landing_url="https://example.test/bad", link_pattern="portfolio"),
        "Example Mutual Fund": static_entry(),
    })
    url = "https://example.test/portfolio-30-Sep-26.xls"
    mock_http({"https://example.test/bad": httpx.Response(500),
               "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
               url: httpx.Response(200, content=workbook())})
    with caplog.at_level(logging.WARNING):
        result = asyncio.run(client.refresh_portfolio_disclosures(db_session))
    assert result.amcs_failed == 1 and result.amcs_processed == 1 and result.schemes_written == 1
    assert "amc=Bad Mutual Fund reason=error" in caplog.text
    assert db_session.query(DisclosureBatch).one().amc_name == "Example Mutual Fund"


def test_ambiguous_sheet_name_logged_unmatched_without_writing(db_session, monkeypatch, mock_http, caplog):
    seed(db_session, base="Example Equity Fund")
    seed(db_session, base="Example Equity Fund Series Two")
    url = "https://example.test/portfolio-30-Sep-26.xlsx"
    with caplog.at_level(logging.INFO):
        result, _ = run_refresh(db_session, monkeypatch, mock_http, {
            "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
            url: httpx.Response(200, content=workbook([("Example Equity", 100, 100)])),
        })
    assert result.schemes_written == 0 and db_session.query(DisclosureBatch).count() == 0
    assert "unmatched" in caplog.text and "Example Equity" in caplog.text
    assert "PORTFOLIO_DISCLOSURE_ALERT" not in caplog.text  # unmatched alerts belong to Task 14


def test_fuzzy_identity_veto_remains_unmatched_through_refresh(db_session, monkeypatch, mock_http, caplog):
    seed(db_session, base="ICICI Prudential Gold ETF FOF")
    url = "https://example.test/portfolio-30-Sep-26.xlsx"
    with caplog.at_level(logging.INFO):
        result, _ = run_refresh(db_session, monkeypatch, mock_http, {
            "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
            url: httpx.Response(200, content=workbook([("ICICI Prudential Silver ETF FOF", 100, 100)])),
        })
    assert result.schemes_written == 0 and db_session.query(DisclosureBatch).count() == 0
    assert "unmatched" in caplog.text


def test_per_amc_commit_and_rollback_preserve_neighbors(db_session, monkeypatch, mock_http):
    amcs = ("First Mutual Fund", "Bad Mutual Fund", "Last Mutual Fund")
    for amc in amcs:
        seed(db_session, amc=amc)
    monkeypatch.setattr(client, "AMC_DISCLOSURE_RESOLVERS", {
        amc: Entry(kind=Kind.STATIC_REGEX, landing_url=f"https://example.test/{index}/list", link_pattern="portfolio")
        for index, amc in enumerate(amcs)
    })
    routes = {}
    for index in range(3):
        url = f"https://example.test/{index}/portfolio-30-Sep-26.xls"
        routes[f"https://example.test/{index}/list"] = httpx.Response(200, text=f'<a href="{url}">Portfolio</a>')
        routes[url] = httpx.Response(200, content=workbook())
    mock_http(routes)
    real_write = client.write_disclosure_batch
    def failing_write(*args, **kwargs):
        count = real_write(*args, **kwargs)
        if args[3] == "Bad Mutual Fund":
            raise RuntimeError("Injected failure after rows were flushed")
        return count
    monkeypatch.setattr(client, "write_disclosure_batch", failing_write)
    result = asyncio.run(client.refresh_portfolio_disclosures(db_session))
    db_session.rollback()  # committed successes must survive an additional caller rollback
    assert result.amcs_processed == 2 and result.amcs_failed == 1 and result.schemes_written == 2
    assert {b.amc_name for b in db_session.query(DisclosureBatch)} == {amcs[0], amcs[2]}
    assert db_session.query(SchemeHolding).count() == 2


def test_wayback_dispatch_runs_through_refresh(db_session, monkeypatch, mock_http):
    seed(db_session)
    entry = Entry(kind=Kind.WAYBACK_LOCATE, landing_url="https://blocked.test/list", cdn_domain="files.test")
    def cdx_response(request):
        assert request.url.params["url"] == entry.landing_url
        return httpx.Response(200, json=cdx_rows([entry.landing_url]))
    real_client = httpx.AsyncClient
    def handler(request):
        if request.url.path == "/cdx/search/cdx":
            return cdx_response(request)
        if request.url.host == "web.archive.org":
            return httpx.Response(200, text='<a href="/portfolio-30-Sep-26.xlsx">Example Equity Fund</a>')
        assert str(request.url) == "https://files.test/portfolio-30-Sep-26.xlsx"
        return httpx.Response(200, content=workbook([("Sheet", 100, 100)]))
    monkeypatch.setattr(client.httpx, "AsyncClient", lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw))
    monkeypatch.setattr(client, "AMC_DISCLOSURE_RESOLVERS", {"Example Mutual Fund": entry})
    result = asyncio.run(client.refresh_portfolio_disclosures(db_session))
    assert result.schemes_written == 1
    assert db_session.query(DisclosureBatch).one().status is DisclosureBatchStatus.WRITTEN


def test_unreadable_standalone_workbook_is_amc_failure(db_session, monkeypatch, mock_http, caplog):
    seed(db_session)
    # The package sniffs as OOXML but the workbook XML cannot be parsed.
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("[Content_Types].xml", b"invalid xml")
        archive.writestr("xl/workbook.xml", b"invalid xml")
    url = "https://example.test/portfolio-30-Sep-26.xlsx"
    with caplog.at_level(logging.WARNING):
        result, _ = run_refresh(db_session, monkeypatch, mock_http, {
            "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
            url: httpx.Response(200, content=buf.getvalue()),
        })
    assert result.amcs_failed == 1 and result.amcs_processed == 0
    assert "amc=Example Mutual Fund reason=error" in caplog.text
    assert db_session.query(DisclosureBatch).count() == 0


def test_crc_bad_zip_member_does_not_stop_valid_peer(db_session, monkeypatch, mock_http, caplog):
    seed(db_session)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_STORED) as archive:
        archive.writestr("bad.xlsx", b"unique corrupt workbook bytes")
        archive.writestr("valid.xls", workbook())
    data = buf.getvalue().replace(b"unique corrupt workbook bytes", b"UNIQUE corrupt workbook bytes")
    url = "https://example.test/portfolio-30-Sep-26.zip"
    with caplog.at_level(logging.WARNING):
        result, _ = run_refresh(db_session, monkeypatch, mock_http, {
            "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
            url: httpx.Response(200, content=data),
        }, static_entry(Shape.ZIP_OF_WORKBOOKS))
    assert result.amcs_processed == 1 and result.schemes_written == 1
    assert db_session.query(DisclosureBatch).one().status is DisclosureBatchStatus.WRITTEN
    assert "amc=Example Mutual Fund reason=parse_error" in caplog.text
    assert "bad.xlsx" in caplog.text


def test_bad_zip_member_alert_names_amc_not_source_url(db_session, monkeypatch, mock_http, caplog):
    seed(db_session)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("bad.xlsx", b"invalid")
        archive.writestr("valid.xlsx", workbook())
    url = "https://example.test/portfolio-30-Sep-26.zip"
    with caplog.at_level(logging.WARNING):
        run_refresh(db_session, monkeypatch, mock_http, {
            "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
            url: httpx.Response(200, content=buf.getvalue()),
        })
    assert "amc=Example Mutual Fund reason=parse_error" in caplog.text
    assert "amc=https://" not in caplog.text


@pytest.mark.parametrize("include_bad_member", [False, True])
def test_zip_without_readable_workbooks_is_amc_failure(db_session, monkeypatch, mock_http, include_bad_member):
    seed(db_session)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as archive:
        archive.writestr("readme.txt", b"no portfolio")
        if include_bad_member:
            archive.writestr("bad.xlsx", b"not a workbook")
    url = "https://example.test/portfolio-30-Sep-26.zip"
    result, _ = run_refresh(db_session, monkeypatch, mock_http, {
        "https://example.test/list": httpx.Response(200, text=f'<a href="{url}">Portfolio</a>'),
        url: httpx.Response(200, content=buf.getvalue()),
    })
    assert result.amcs_failed == 1 and result.amcs_processed == 0
    assert db_session.query(DisclosureBatch).count() == 0


def test_canara_real_filename_hint_writes_code_named_fixture(db_session, monkeypatch):
    seed(db_session, amc="Canara Robeco Mutual Fund", base="Canara Robeco Small Cap Fund")
    entry = AMC_DISCLOSURE_RESOLVERS["Canara Robeco Mutual Fund"]
    url = "https://www.canararobeco.com/wp-content/uploads/2026/10/SC-%E2%80%93-Canara-Robeco-Small-Cap-Fund-%E2%80%93-September-2026.xlsx"
    fixture = Path(__file__).resolve().parents[2] / "fixtures/portfolio_disclosures/canara-robeco-small-cap_asof_2026-09-30.xlsx"
    data = fixture.read_bytes()
    real_client = httpx.AsyncClient
    def handler(request):
        if request.url.host == "web.archive.org":
            assert request.url.params["matchType"] == "prefix"
            return httpx.Response(200, json=cdx_rows([url]))
        assert str(request.url) == url
        return httpx.Response(200, content=data)
    monkeypatch.setattr(client.httpx, "AsyncClient", lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw))
    monkeypatch.setattr(client, "AMC_DISCLOSURE_RESOLVERS", {"Canara Robeco Mutual Fund": entry})
    result = asyncio.run(client.refresh_portfolio_disclosures(db_session))
    assert result.schemes_written == 1
    batch = db_session.query(DisclosureBatch).one()
    assert batch.status is DisclosureBatchStatus.WRITTEN
    assert batch.as_of_date == date(2026, 9, 30)
    assert batch.source_url == url and db_session.query(SchemeHolding).count() > 1


def locate_run3(amc, handler, monkeypatch, as_of=date(2026, 9, 30)):
    monkeypatch.setattr(client, "_latest_month_end", lambda: as_of)
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler),
                                    headers=client._HEADERS, follow_redirects=True) as http:
            return await client._locate_and_fetch(http, AMC_DISCLOSURE_RESOLVERS[f"{amc} Mutual Fund"], [])
    return asyncio.run(run())


def assert_located(files, urls, hints=None, as_of=date(2026, 9, 30)):
    assert [f.source_url for f in files] == urls
    assert all(f.as_of_date == as_of and f.data == b"PK\x03\x04fixture" for f in files)
    if hints is not None:
        assert [f.scheme_hint for f in files] == hints


def test_run3_franklin_sitemap_and_title_date_not_reference_date(monkeypatch):
    # Capture log §12 Franklin: document[], contentGrouping/dctermsTitle/downloadUrl.
    url = "https://franklintempletonprod.widen.net/content/captured/original/Franklin_India_Long_Term_Fund_Sep_26_Monthly_Portfolio.xls?u=pl55mb&download=true"
    seen = []
    def handler(r):
        seen.append(str(r.url))
        if r.url.path == "/sitemap.xml":
            return httpx.Response(200, text='<urlset><url><loc>https://www.franklintempletonindia.com/fund-details/fund-overview/42352</loc></url><url><loc>https://www.franklintempletonindia.com/fund-details/fund-overview/42352/reports</loc></url></urlset>')
        if r.url.path == "/api/literature/v1/documents":
            assert r.method == "GET" and dict(r.url.params) == {"channel": "en-in", "cache": "true", "frkOneTISNumber": "42352"}
            return httpx.Response(200, json={"document": [
                {"contentGrouping": "Dis-Portfolio-Monthly", "dctermsTitle": "As on September 30, 2026 (Uploaded on October 08, 2026)", "frkReferenceDate": "2026-10-08", "downloadUrl": url},
                {"contentGrouping": "Dis-Portfolio-Fortnightly", "dctermsTitle": "As on September 30, 2026", "downloadUrl": "https://bad.test/fortnightly.xls"},
                {"contentGrouping": "Dis-Portfolio-Monthly", "dctermsTitle": "As on August 31, 2026", "frkReferenceDate": "2026-09-09", "downloadUrl": "https://bad.test/old.xls"},
            ]})
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Franklin Templeton", handler, monkeypatch), [url], ["Franklin India Long Term Fund"])
    assert len(seen) == 3


@pytest.mark.parametrize("token_data", [{"token": "Bearer captured-token"}, "{'token': 'Bearer captured-token'}"])
def test_run3_axis_token_chain_and_misspelled_url_key(monkeypatch, token_data):
    # Capture log Axis: data string/dict, data.documentList[].docuementURL.
    url = "https://www.axismf.com/1/5/464/560/3622/4593/Monthly_Portfolio_30_09_2026_75011dd855.xlsx"
    def handler(r):
        if r.url.path == "/cms/token":
            assert r.method == "POST" and json.loads(r.content) == {}
            return httpx.Response(200, json={"status": "success", "statusCode": "0", "data": token_data})
        if r.url.path == "/cms/get-scheme-documents":
            assert r.headers["Authorization"] == "Bearer captured-token"
            assert json.loads(r.content) == {"sdType": "yearMonthSchemeDocs", "sdID": "sdMonthSchemePortfolio", "year": "2026", "month": "September", "schemeCode": "Consolidated"}
            return httpx.Response(200, json={"data": {"documentList": [
                {"documentName": "Weekly Debt Portfolio 30-09-2026", "docuementURL": "https://bad.test/weekly.xlsx"},
                {"documentName": "Monthly Portfolio 30-09-2026", "docuementURL": url, "documentPostedDate": "2026-10-09"},
            ]}})
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Axis", handler, monkeypatch), [url], [None])


def test_run3_lic_consolidated_form_chain_returns_both_files(monkeypatch):
    # Capture log LIC consolidated route: option fragments; two equity/debt anchors.
    urls = ["https://www.licmf.com/assets/pdf/statuary_disclosure_new/21/LIC%20MF%20Monthly%20" + name + "%20Portfolio_Sep_26.xlsx" for name in ("Equity", "Debt")]
    def handler(r):
        body = dict(httpx.QueryParams(r.content.decode()))
        if r.url.path == "/downloads/consolidated-portfolio-filters":
            assert r.method == "POST" and r.headers["X-Requested-With"] == "XMLHttpRequest"
            assert r.headers["Content-Type"].startswith("application/x-www-form-urlencoded")
            if body == {"id": "639", "filter": "year"}:
                return httpx.Response(200, text="<option value='2026'>2026</option>")
            assert body == {"year": "2026", "id": "639", "filter": "month"}
            return httpx.Response(200, text="<option value='9'>September</option>")
        if r.url.path == "/downloads/consolidated-portfolio-files":
            assert body == {"id": "639", "month": "9", "year": "2026"}
            return httpx.Response(200, text='<a href="/assets/pdf/statuary_disclosure_new/21/LIC MF Monthly Equity Portfolio_Sep_26.xlsx">Monthly Portfolio Equity as on September 30, 2026</a><a href="/assets/pdf/statuary_disclosure_new/21/LIC MF Monthly Debt Portfolio_Sep_26.xlsx">Monthly Portfolio Debt as on September 30, 2026</a>')
        assert str(r.url) in urls
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("LIC", handler, monkeypatch), urls, [None, None])


def test_run3_sbi_html_fragment_prefers_all_schemes(monkeypatch):
    # Capture log SBI: GetSchemePortfolioSheets (without QS), HTML response.
    url = "https://www.sbimf.com/docs/default-source/scheme-portfolios/all-schemes-monthly-portfolio---as-on-30th-september-2026.xlsx?sfvrsn=2a79126c_2"
    def handler(r):
        if r.url.path == "/ajaxcall/CMS/GetSchemePortfolioSheets":
            assert r.method == "POST" and r.headers["X-Requested-With"] == "XMLHttpRequest"
            assert json.loads(r.content) == {"FundId": "", "PSYear": "2026", "PSMonth": "September", "PSFrequency": "Monthly"}
            return httpx.Response(200, text=f'<table><a href="/docs/default-source/scheme-portfolios/sbi-balanced-advantage-fund-monthly-portfolio---september-2026.xlsx">Portfolio</a><a href="{url}">All Schemes Monthly Portfolio</a></table>')
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("SBI", handler, monkeypatch), [url], [None])


def test_run3_trust_title_dates_opaque_url_and_follows_apex_redirect(monkeypatch):
    # Capture log Trust: resultSetArray[].fileurl; filename is upload timestamp.
    apex = "https://trustmf.com/Content/2026/10/Monthly%20Port_20261009132916.xlsx"
    live = apex.replace("https://trustmf.com", "https://www.trustmf.com")
    def handler(r):
        if r.url.path == "/api/api/Trust/GetData":
            assert json.loads(r.content) == {"systemQueryFileName": "disclosuresweb.xml", "tagName": "GetDisclosureByType", "searchField": "", "searchValue": "", "sortField": "uploaddate", "sortDirection": "DESC", "replaceField": "_slug_", "replaceValue": "portfolio-monthly-disclosure"}
            return httpx.Response(200, json={"resultSetArray": [{"title": "TRUSTMF Monthly Portfolio Report as on 30.09.2026", "uploaddate": "10/9/2026 12:00:00 AM", "fileurl": apex.replace("%20", " ")}]})
        if str(r.url) == apex:
            return httpx.Response(307, headers={"Location": live})
        assert str(r.url) == live
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Trust", handler, monkeypatch), [apex], [None])


def test_run3_mirae_title_scheme_hint_and_redirected_html_rejected(monkeypatch):
    # Capture log Mirae: Data[].URL/Title, DataCount, 302 soft-404.
    url = "https://www.miraeassetmf.co.in/docs/default-source/portfolios/mafcf_sept2026.xlsx"
    def handler(r):
        if r.url.path == "/AjaxService/GetDownloadsData":
            assert json.loads(r.content) == {"request": {"modulename": "portfolio_tab1", "pgno": 1, "pgsize": 200}}
            return httpx.Response(200, json={"Data": [
                {"Title": "Portfolio Details as on 30th September 2026 for Mirae Asset Flexi Cap Fund", "URL": "/docs/default-source/portfolios/mafcf_sept2026.xlsx"},
                {"Title": "Portfolio Details as on 30th September 2026 for Mirae Asset Missing Fund", "URL": "/docs/default-source/portfolios/missing_sept2026.xlsx"},
            ], "DataCount": 2})
        if "missing_" in r.url.path:
            return httpx.Response(302, headers={"Location": "/error?aspxerrorpath=missing"})
        if r.url.path == "/error":
            return httpx.Response(200, text="<html>Not found</html>")
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Mirae Asset", handler, monkeypatch), [url], ["Mirae Asset Flexi Cap Fund"])


def test_run3_boi_case_sensitive_folder_and_nested_json_string(monkeypatch):
    # Capture log Bank of India: d is JSON text; Documents[].FolderUrl/DocName.
    url = "https://www.boimf.in/docs/default-source/investorcorner/monthly-portfolio/monthly-portfolio---30-sep-2026.xlsx?sfvrsn=8dbd0a0e_3"
    def handler(r):
        if r.url.path == "/AjaxService.asmx/GetDocuments":
            assert r.url.host == "www.boimf.in" and r.method == "POST"
            assert json.loads(r.content) == {"pagno": 0, "category": None, "fromDate": None, "toDate": None, "LibraryName": "InvestorCorner", "folderName": "MONTHLY PORTFOLIO", "CategoryValue": None}
            return httpx.Response(200, json={"d": json.dumps({"Documents": [{"DocName": "MONTHLY-PORTFOLIO - 30-SEP-2026", "FolderUrl": url}]})})
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Bank of India", handler, monkeypatch), [url], [None])


@pytest.mark.parametrize("as_of,month_id,has_files", [(date(2026, 9, 30), "436", False), (date(2026, 3, 31), "435", True)])
def test_run3_taurus_select_ids_half_yearly_and_unpublished_september(monkeypatch, as_of, month_id, has_files):
    # Capture log Taurus: year 562, months 435/436; March href starts with space.
    url = "https://www.taurusmutualfund.com/sites/default/files/downloads/Taurus_Flexi_Cap_Fund_Half_Yearly_Portfolio_Report_Performance_March_2026.xlsx"
    def handler(r):
        if r.url.path == "/financials/portfolio":
            if not r.url.query:
                return httpx.Response(200, text='<select name="field_hy_portfolio_year_target_id"><option value="562">2026</option></select><select name="field_hy_portfolio_month_target_id"><option value="435">March</option><option value="436">September</option></select>')
            assert dict(r.url.params) == {"field_hy_portfolio_year_target_id": "562", "field_hy_portfolio_month_target_id": month_id}
            return httpx.Response(200, text=f'<a href=" /sites/default/files/downloads/Taurus_Flexi_Cap_Fund_Half_Yearly_Portfolio_Report_Performance_March_2026.xlsx">Portfolio</a>' if has_files else "<html>No file links</html>")
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    files = locate_run3("Taurus", handler, monkeypatch, as_of)
    if has_files:
        assert_located(files, [url], ["Taurus Flexi Cap Fund"], as_of)
    else:
        assert files == []


def test_run3_absl_captured_accordion_request_and_dead_host_rewrite(monkeypatch):
    # Capture log ABSL: AccordionList[].pdfUrl/ResourceLink; varying ZIP stems.
    dead = "https://abcscprod.azureedge.net/-/media/bsl/files/resources/monthly-portfolio/2026/30092026_abslmf_monthlydisclosure.zip"
    live = dead.replace("abcscprod.azureedge.net", "mutualfund.adityabirlacapital.com")
    def handler(r):
        if r.url.path.endswith("/FactsheetAccordionById"):
            assert r.method == "GET" and dict(r.url.params) == {"id": "3ccab227-9de5-4494-b78d-2b4f7c0c054a", "ctype": "/sitecore/content/Root/BSL/Library/Lists/FAQ/Customer Types/Individual", "month": " ", "year": "0"}
            return httpx.Response(200, json={"AccordionList": [{"ResourceLink": "Monthly Portfolios as on September 30, 2026", "IconClass": "zip", "pdfUrl": dead}]})
        assert str(r.url) == live
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Aditya Birla Sun Life", handler, monkeypatch), [live], [None])


@pytest.mark.parametrize("as_of,fy,month", [(date(2026, 9, 30), "2026-2027", "September"), (date(2026, 1, 31), "2025-2026", "January")])
def test_run3_navi_nonce_header_category_and_financial_year(monkeypatch, as_of, fy, month):
    # Capture log Navi: navi_property rest_url/nonce; data[].title/url, form POST.
    title = f"Navi Flexi Cap Fund 1st &#8211; {as_of.day}th {month} {as_of.year}"
    url = f"https://public-assets.prod.navi-tech.in/navi-website-assests/documents/Navi_Flexi_Cap_Fund_1st_{as_of.day}th_{month}_{as_of.year}_20261009174229.xlsx"
    def handler(r):
        if r.url.path == "/mutual-fund/downloads/portfolio":
            return httpx.Response(200, text='<script>navi_property = {"ajaxurl":"/ajax","rest_url":"https://navi.com/wp-json/","nonce":"abc123","filter_nonce":"wrong"}</script>')
        if r.url.path == "/wp-json/nv/v1/documents":
            assert r.method == "POST" and r.headers["WP-NONCE"] == "abc123"
            assert dict(httpx.QueryParams(r.content.decode())) == {"financial_year": fy, "value": month, "category": "884", "type": "Monthly", "order": "DESC"}
            return httpx.Response(200, json={"success": True, "data": [{"title": title, "url": url}]})
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Navi", handler, monkeypatch, as_of), [url], ["Navi Flexi Cap Fund"], as_of)


def test_run3_pgim_nested_tabs_and_report_date(monkeypatch):
    # Capture log PGIM: data[].content[].pdfPath/title/dateMonthYear/month/year.
    url = "https://www.pgimindia.com/api/v1/brochure/about-us/image/PGIM%20INDIA%20FLEXI%20CAP%20FUND%20Sep%202026.xlsx"
    def handler(r):
        if r.url.path == "/api/v1/brochure/published/disclosure":
            assert json.loads(r.content) == {"headerId": 2, "sectionId": "SECTION_747960037", "source": "W", "branchCode": None}
            return httpx.Response(200, json={"resultInfo": {}, "data": [{"tabName": "Equity", "content": [
                {"title": "PGIM INDIA FLEXI CAP FUND Sep 2026", "pdfPath": url.replace("%20", " "), "dateMonthYear": "30 September 2026", "month": "September", "year": "2026"},
                {"title": "PGIM INDIA FLEXI CAP FUND Aug 2026", "pdfPath": "https://bad.test/old.xlsx", "dateMonthYear": "31 August 2026", "month": "August", "year": "2026"},
            ]}]})
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("PGIM India", handler, monkeypatch), [url], ["PGIM INDIA FLEXI CAP FUND"])


def test_run3_whiteoak_strapi_paging_monthly_filter_and_project_ua(monkeypatch):
    # Capture log WhiteOak: data[].attributes.doc_file.data.attributes.url;
    # meta.pagination, published_date is upload date, not disclosure date.
    url = "https://content.whiteoakamc.com/WOC_Special_Opportunities_Fund_Sept_2026_e22c53bbd6.xlsx"
    def handler(r):
        assert r.headers["User-Agent"] == client._HEADERS["User-Agent"]
        if r.url.path == "/api/scheme-portfolios":
            p = dict(r.url.params)
            assert p == {"pagination[page]": p["pagination[page]"], "pagination[pageSize]": "100", "sort[0]": "published_date:desc", "sort[1]": "id:desc", "populate": "*"}
            attrs = {"scheme_name": "WhiteOak Capital Special Opportunities Fund", "period": "Monthly", "doc_name": "WhiteOak Capital Special Opportunities Fund Monthly Portfolio Disclosure - 30th September2026", "doc_link": None, "published_date": "2026-10-09", "doc_file": {"data": {"attributes": {"url": url}}}}
            if p["pagination[page]"] == "1":
                attrs = dict(attrs, period="Fortnightly")
            return httpx.Response(200, json={"data": [{"id": 1560, "attributes": attrs}], "meta": {"pagination": {"page": int(p["pagination[page]"]), "pageCount": 2, "pageSize": 100, "total": 2}}})
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("WhiteOak Capital", handler, monkeypatch), [url], ["WhiteOak Capital Special Opportunities Fund"])


@pytest.mark.parametrize("blocked", [False, True])
def test_run3_icici_complete_headers_json_gate_and_blob_path(monkeypatch, blocked):
    # Capture log ICICI: success.data.files[].title.text/url; F5 HTML gate.
    url = "https://www.icicipruamc.com/blob/downloads/Files/Monthly%20Portfolio%20Disclosures/2026/Sep/Monthly-Portfolio-Disclosure-September-2026.zip"
    def handler(r):
        if r.url.host == "apimf.icicipruamc.com":
            assert r.method == "POST" and r.url.path == "/nms/v1/downloads/files"
            assert r.headers["sourceurl"] == "DOWNLOADS" and r.headers["env"] == "api"
            uuid.UUID(r.headers["requestapiid"])
            assert r.headers["Origin"] == "https://www.icicipruamc.com"
            assert r.headers["Referer"].startswith("https://www.icicipruamc.com/media-center/downloads")
            assert "application/json" in r.headers["Accept"]
            assert json.loads(r.content) == {"categoryId": "26a073d7-08d2-4a95-95fa-f83a4ee51e40", "schemeCategory": "", "userType": "Investor", "fileType": "All", "page": "1", "size": "20", "filter": [], "categoryName": "OTHERS"}
            if blocked:
                return httpx.Response(200, text="<html>bobcmn /TSPD/ challenge</html>", headers={"Content-Type": "text/html"})
            return httpx.Response(200, json={"success": {"data": {"files": [{"title": {"text": "Monthly Portfolio Disclosure September 2026"}, "url": "/downloads/Files/Monthly Portfolio Disclosures/2026/Sep/Monthly-Portfolio-Disclosure-September-2026.zip"}], "isNext": False}}})
        assert not blocked and str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    if blocked:
        with pytest.raises(ValueError, match="BLOCKED.*JSON"):
            locate_run3("ICICI Prudential", handler, monkeypatch)
    else:
        assert_located(locate_run3("ICICI Prudential", handler, monkeypatch), [url], [None])


def test_run3_uti_exact_doc_zip_and_no_invented_scheme_hint(monkeypatch):
    # Capture log UTI: rows[0].doc; raw codes map only through an exact name source.
    url = "https://d3ce1o48hc5oli.cloudfront.net/static/generic-zip/Oct-26/FW__UTI_MF_SCHEME_PORTFOLIOS_-_30.09.2026_.zip"
    def handler(r):
        if r.url.path == "/api/get-consolidate-portfolio-disclosure":
            assert dict(r.url.params) == {"year": "2026", "month": "September"}
            return httpx.Response(200, json={"rows": [{"name": "Consolidated Portfolio September 2026", "type": "Zip", "doc": url, "url": url, "category": "Consolidate portfolio disclosure", "month": "September", "year": "2026"}]})
        assert str(r.url) == url
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("UTI", handler, monkeypatch), [url], [None])


@pytest.mark.parametrize("method", ["GET", "POST"])
def test_run3_generic_json_path_regex_placeholders_and_rewrite(method):
    entry = Entry(kind=Kind.JSON_API, endpoint_url="https://example.test/api", http_method=method,
                  request_params={"year": "{yyyy}", "month": "{mm}"}, request_body={"month": "{month_name}"},
                  response_json_path="data.files", documents=(r"Monthly Portfolio {month_name} {yyyy}",),
                  rewrite_host=("abcscprod.azureedge.net", "mutualfund.adityabirlacapital.com"))
    live = "https://mutualfund.adityabirlacapital.com/-/media/portfolio-30-Sep-26.zip"
    def handler(r):
        if r.url.path == "/api":
            assert r.method == method and dict(r.url.params) == {"year": "2026", "month": "09"}
            if method == "POST":
                assert json.loads(r.content) == {"month": "September"}
            return httpx.Response(200, json={"data": {"files": [
                {"name": "Monthly Portfolio September 2026", "url": live.replace("mutualfund.adityabirlacapital.com", "abcscprod.azureedge.net")},
                {"name": "Other Document September 2026", "url": "https://bad.test/notselected.zip"},
            ]}})
        assert str(r.url) == live
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_json_api(http, entry, date(2026, 9, 30))
    assert_located(asyncio.run(run()), [live], [None])


@pytest.mark.parametrize("rows", [[], [{"name": "Portfolio", "url": "https://bad.test/current.zip"}]])
def test_run3_generic_empty_or_undated_response_has_no_download(rows):
    entry = Entry(kind=Kind.JSON_API, endpoint_url="https://example.test/api", response_json_path="data", documents=())
    def handler(r):
        assert r.url.path == "/api"
        return httpx.Response(200, json={"data": rows})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_json_api(http, entry, date(2026, 9, 30))
    assert asyncio.run(run()) == []


@pytest.mark.parametrize("amc", ["Franklin Templeton", "Axis", "LIC", "SBI", "Trust", "Mirae Asset",
                                "Bank of India", "Taurus", "Aditya Birla Sun Life", "Navi", "PGIM India",
                                "WhiteOak Capital", "ICICI Prudential", "UTI"])
def test_run3_every_scoped_locator_empty_response_returns_no_files(monkeypatch, amc):
    # Same capture-log response containers, trimmed to no published documents.
    def handler(r):
        assert not client._WORKBOOK_URL.search(str(r.url)), "Empty response must not download a file"
        if r.url.path == "/sitemap.xml":
            return httpx.Response(200, text='<loc>https://www.franklintempletonindia.com/fund-details/fund-overview/42352</loc>')
        if r.url.path == "/cms/token":
            return httpx.Response(200, json={"data": {"token": "Bearer captured-token"}})
        if r.url.path == "/cms/get-scheme-documents":
            return httpx.Response(200, json={"data": {"documentList": []}})
        if amc in {"LIC", "SBI", "Taurus"}:
            return httpx.Response(200, text="<html></html>")
        if r.url.path == "/mutual-fund/downloads/portfolio":
            return httpx.Response(200, text='<script>navi_property={"rest_url":"https://navi.com/wp-json/","nonce":"abc123"}</script>')
        payloads = {
            "Franklin Templeton": {"document": []}, "Trust": {"resultSetArray": []},
            "Mirae Asset": {"Data": [], "DataCount": 0}, "Bank of India": {"d": '{"Documents":[]}'},
            "Aditya Birla Sun Life": {"AccordionList": []}, "Navi": {"success": True, "data": []},
            "PGIM India": {"data": []}, "WhiteOak Capital": {"data": [], "meta": {"pagination": {"pageCount": 1}}},
            "ICICI Prudential": {"success": {"data": {"files": [], "isNext": False}}}, "UTI": {"rows": []},
        }
        return httpx.Response(200, json=payloads[amc])
    assert locate_run3(amc, handler, monkeypatch) == []


def test_run3_mirae_paging_keeps_all_target_scheme_files(monkeypatch):
    # Capture log Mirae DataCount/pgno: two pages, older rows never downloaded.
    first = "https://www.miraeassetmf.co.in/docs/default-source/portfolios/mafcf_sept2026.xlsx"
    second = "https://www.miraeassetmf.co.in/docs/default-source/portfolios/macmf_sept2026.xlsx"
    def handler(r):
        if r.url.path == "/AjaxService/GetDownloadsData":
            page = json.loads(r.content)["request"]["pgno"]
            row = {"Title": "Portfolio Details as on 30th September 2026 for Mirae Asset Flexi Cap Fund", "URL": first}
            if page == 1:
                rows = [row] * 200
            else:
                assert page == 2
                rows = [{"Title": "Portfolio Details as on 30th September 2026 for Mirae Asset Liquid Fund", "URL": second}]
            return httpx.Response(200, json={"Data": rows, "DataCount": 201})
        assert str(r.url) in {first, second}
        return httpx.Response(200, content=b"PK\x03\x04fixture")
    assert_located(locate_run3("Mirae Asset", handler, monkeypatch), [first, second], ["Mirae Asset Flexi Cap Fund", "Mirae Asset Liquid Fund"])


def test_run3_uti_real_zip_code_blocks_remain_unmatched(db_session, monkeypatch, mock_http, caplog):
    # Capture log UTI: exact SCHEME: source exists but no code-to-name mapping is
    # wired in this run. Never assign every EXPOSURE block to one family.
    seed(db_session, amc="UTI Mutual Fund", base="UTI - Unit Linked Insurance Plan")
    fixture = Path(__file__).resolve().parents[2] / 'fixtures/portfolio_disclosures/uti-consolidated_asof_2026-09-30.zip'
    url = "https://d3ce1o48hc5oli.cloudfront.net/static/generic-zip/Oct-26/FW__UTI_MF_SCHEME_PORTFOLIOS_-_30.09.2026_.zip"
    api = "https://www.utimf.com/api/get-consolidate-portfolio-disclosure?year=2026&month=September"
    monkeypatch.setattr(client, "_latest_month_end", lambda: date(2026, 9, 30))
    with caplog.at_level(logging.INFO):
        result, _ = run_refresh(db_session, monkeypatch, mock_http, {
            api: httpx.Response(200, json={"rows": [{"name": "Consolidated Portfolio September 2026", "doc": url}]}),
            url: httpx.Response(200, content=fixture.read_bytes()),
        }, AMC_DISCLOSURE_RESOLVERS["UTI Mutual Fund"], amc="UTI Mutual Fund")
    assert result.amcs_processed == 1 and result.schemes_written == 0
    assert db_session.query(DisclosureBatch).count() == 0
    assert "unmatched" in caplog.text and "EXPOSURE:002" in caplog.text


def test_run3_upload_only_title_is_not_a_disclosure_date():
    # Franklin explicitly distinguishes as-on from uploaded-on in dctermsTitle.
    assert client._disclosed_date("Monthly Portfolio (Uploaded on September 30, 2026)") is None


def test_run3_whiteoak_monthly_row_without_attached_file_is_skipped(monkeypatch):
    # Strapi nullable relations must not turn an unavailable file into today's data.
    def handler(r):
        assert r.url.path == "/api/scheme-portfolios"
        return httpx.Response(200, json={"data": [{"attributes": {"period": "Monthly", "scheme_name": "WhiteOak Capital Liquid Fund", "doc_name": "Monthly Portfolio Disclosure - 30th September2026", "doc_file": None}}], "meta": {"pagination": {"pageCount": 1}}})
    assert locate_run3("WhiteOak Capital", handler, monkeypatch) == []


def test_html_served_under_a_workbook_name_is_not_a_file():
    # Mirae-style: HTTP 200 with an HTML error page under a .xlsx name (F18).
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r: httpx.Response(
                200, content=b"<html>Not found</html>"))) as http:
            return await client._download_workbook(http, "https://example.test/x.xlsx")
    assert asyncio.run(run()) is None


def test_wayback_prefix_search_ignores_dated_pdfs_and_images_that_would_win_on_date():
    entry = Entry(kind=Kind.WAYBACK_LOCATE, landing_url="https://www.example.test/wp-content/uploads/",
                  link_pattern=r"Example-Equity-Fund")
    xlsx = "https://www.example.test/wp-content/uploads/2026/10/Example-Equity-Fund-September-2026.xlsx"
    pdf = "https://www.example.test/wp-content/uploads/2026/10/Example-Equity-Fund-SID-October-2026.pdf"
    def handler(request):
        if request.url.host == "web.archive.org":
            return httpx.Response(200, json=cdx_rows([xlsx, pdf]))
        assert str(request.url) == xlsx  # the PDF must never be requested
        return httpx.Response(200, content=workbook())
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as http:
            return await client._locate_wayback(http, entry)
    assert [f.source_url for f in asyncio.run(run())] == [xlsx]


def test_empty_204_json_response_means_no_documents_not_an_error():
    # Franklin answers HTTP 204 with an empty body for funds that have no documents
    # (found live 2026-10-13: 6 of 52 funds); one such fund must not fail the whole AMC.
    req = httpx.Request("GET", "https://example.test/api")
    assert client._response_json(httpx.Response(204, headers={"Content-Type": "application/json"}, request=req)) == {}
    assert client._response_json(httpx.Response(200, headers={"Content-Type": "application/json"}, content=b"", request=req)) == {}
