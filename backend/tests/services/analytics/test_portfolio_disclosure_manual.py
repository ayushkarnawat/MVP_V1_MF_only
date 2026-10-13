import io
import uuid
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from app.models.enums import DisclosureBatchStatus as Status
from app.models.reference import DisclosureBatch, Scheme, SchemeHolding
from app.services.analytics import portfolio_disclosure_manual as manual


def seed(db, base='ITI Flexi Cap Fund', amc='ITI Mutual Fund'):
    scheme = Scheme(id=uuid.uuid4(), name=base+' - Direct Growth', base_name=base,
                    amc_name=amc, sebi_category='Equity Scheme - Flexi Cap Fund')
    db.add(scheme)
    db.commit()
    return scheme


def workbook(name='ITI Flexi Cap Fund'):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = name
    ws.append(['Name of the Instrument', 'ISIN', 'Industry', 'Quantity', 'Market Value (Rs. in Lakhs)', '% to NAV'])
    ws.append(['Equity & Equity related'])
    ws.append(['Bank', 'INE040A01034', 'Banks', 10, 10, 100])
    ws.append(['Grand Total', None, None, None, 10, 100])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


def test_queue_idempotent_and_summary_latest_only(db_session):
    scheme = seed(db_session)
    assert manual.ensure_manual_queue_rows(db_session, date(2026, 10, 3)) == 1
    assert manual.ensure_manual_queue_rows(db_session, date(2026, 10, 3)) == 0
    db_session.add(DisclosureBatch(scheme_id=scheme.id, as_of_date=date(2026, 8, 31),
                                  amc_name=scheme.amc_name, status=Status.AWAITING_MANUAL_UPLOAD))
    db_session.flush()
    assert manual.manual_queue_summary(db_session, date(2026, 10, 3)) == {'ITI Mutual Fund': 1}
    result = manual.import_manual_files(db_session, scheme.amc_name,
        [('portfolio-30-Sep-26.xls', workbook())], None, date(2026, 10, 3))
    assert result == {'matched': 1, 'unmatched': 0, 'schemes_written': 1}
    assert manual.manual_queue_summary(db_session, date(2026, 10, 3)) == {}
    assert db_session.query(SchemeHolding).count() == 1
    assert db_session.query(DisclosureBatch).filter_by(as_of_date=date(2026, 9, 30)).one().source_url == 'manual:portfolio-30-Sep-26.xls'


@pytest.mark.parametrize('files,as_of,allow_old', [
    ([('good.xlsx', workbook()), ('bad.xlsx', b'%PDF-bad')], date(2026, 9, 30), False),
    ([('undated.xlsx', workbook())], None, False),
    ([('good.xlsx', workbook())], date(2026, 10, 31), True),
    ([('good.xlsx', workbook())], date(2026, 7, 31), False),
])
def test_invalid_call_writes_nothing(db_session, files, as_of, allow_old):
    seed(db_session)
    with pytest.raises(ValueError):
        manual.import_manual_files(db_session, 'ITI Mutual Fund', files, as_of,
                                   date(2026, 10, 3), allow_old=allow_old)
    assert db_session.query(DisclosureBatch).count() == 0


def test_non_manual_rejected(db_session):
    with pytest.raises(ValueError, match='manual'):
        manual.import_manual_files(db_session, 'Axis Mutual Fund', [], None, date(2026, 10, 3))


def test_old_date_requires_override_and_empty_input_rejected(db_session):
    seed(db_session)
    with pytest.raises(ValueError):
        manual.import_manual_files(db_session, 'ITI Mutual Fund', [], date(2026, 9, 30), date(2026, 10, 3))
    assert manual.import_manual_files(db_session, 'ITI Mutual Fund', [('old.xlsx', workbook())],
        date(2026, 7, 31), date(2026, 10, 3), allow_old=True)['schemes_written'] == 1


@pytest.mark.parametrize('filename,amc,base,as_of', [
    ('edelweiss_asof_2026-09-30.xlsx', 'Edelweiss Mutual Fund', 'Edelweiss Flexi Cap Fund', date(2026, 9, 30)),
    ('iti_asof_2026-08-31.xlsx', 'ITI Mutual Fund', 'ITI Flexi Cap Fund', date(2026, 8, 31)),
    ('jm-financial-large-cap_asof_2026-09-30.xlsx', 'JM Financial Mutual Fund', 'JM Large Cap Fund', date(2026, 9, 30)),
    ('jm-financial-mid-cap_asof_2026-09-30.xlsx', 'JM Financial Mutual Fund', 'JM Midcap Fund', date(2026, 9, 30)),
])
def test_real_manual_fixtures_end_to_end(db_session, filename, amc, base, as_of):
    scheme = seed(db_session, base, amc)
    data = (Path(__file__).parents[2] / 'fixtures/portfolio_disclosures' / filename).read_bytes()
    if amc == 'Edelweiss Mutual Fund':
        # A realistic family pool prevents other disclosed names fuzzily falling
        # into the only seeded family. Names come directly from the real Index.
        from app.services.analytics.portfolio_disclosure_parsing import read_index_sheet
        from app.services.analytics.portfolio_disclosure_matching import _canonical_name
        for name in set(read_index_sheet(data).values()):
            if _canonical_name(name) != _canonical_name(base):
                seed(db_session, name, amc)
    result = manual.import_manual_files(db_session, amc, [(filename, data)], as_of,
                                       date(2026, 10, 13), allow_old=True)
    assert result['matched'] >= 1 and result['schemes_written'] >= 1
    batch = db_session.query(DisclosureBatch).filter_by(scheme_id=scheme.id).one()
    assert batch.status is Status.WRITTEN
    assert batch.as_of_date == as_of and batch.source_url == 'manual:'+filename
    assert db_session.query(SchemeHolding).filter_by(batch_id=batch.id).count() > 0
def test_manual_zip_and_cli_run_commit_queue(db_session,tmp_path):
    import importlib.util
    import zipfile
    scheme=seed(db_session)
    manual.ensure_manual_queue_rows(db_session,date(2026,10,3))
    db_session.commit()
    bundle=io.BytesIO()
    with zipfile.ZipFile(bundle,'w') as archive:
        archive.writestr('holdings.xls',workbook())
    path=tmp_path/'manual-30-Sep-26.zip'
    path.write_bytes(bundle.getvalue())
    script=Path(__file__).parents[3]/'scripts/jobs/import_manual_portfolio_disclosures.py'
    spec=importlib.util.spec_from_file_location('manual_portfolio_job',script)
    job=importlib.util.module_from_spec(spec)
    spec.loader.exec_module(job)
    assert job.run(db_session,'ITI Mutual Fund',[path],None,date(2026,10,3)) == {'matched':1,'unmatched':0,'schemes_written':1}
    db_session.rollback()
    assert db_session.query(DisclosureBatch).filter_by(scheme_id=scheme.id).one().status is Status.WRITTEN


def test_manual_parse_failure_in_second_file_leaves_existing_queue(db_session):
    seed(db_session)
    manual.ensure_manual_queue_rows(db_session,date(2026,10,3))
    db_session.commit()
    bad=workbook()
    wb=openpyxl.load_workbook(io.BytesIO(bad))
    wb.active.cell(1,5,'Market Value (Rs. in Millions)')
    buf=io.BytesIO()
    wb.save(buf)
    wb.close()
    with pytest.raises(ValueError):
        manual.import_manual_files(db_session,'ITI Mutual Fund',
            [('valid.xlsx',workbook()),('invalid.xlsx',buf.getvalue())],date(2026,9,30),date(2026,10,3))
    assert db_session.query(DisclosureBatch).one().status is Status.AWAITING_MANUAL_UPLOAD
    assert db_session.query(SchemeHolding).count()==0
