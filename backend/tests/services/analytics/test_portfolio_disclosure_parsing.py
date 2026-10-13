# backend/tests/services/analytics/test_portfolio_disclosure_parsing.py
import io
from decimal import Decimal

import openpyxl
import pytest
import xlwt

from app.services.analytics.portfolio_disclosure_parsing import (
    find_header_row, market_value_multiplier, parse_sheet_rows, parse_workbook,
    sniff_workbook_format,
)


def _xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xls_bytes(sheets: dict[str, list[list]]) -> bytes:
    wb = xlwt.Workbook()
    for name, rows in sheets.items():
        ws = wb.add_sheet(name)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                ws.write(r, c, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sniff_workbook_format_detects_both():
    assert sniff_workbook_format(_xlsx_bytes({"S1": [["a"]]})) == "xlsx"
    assert sniff_workbook_format(_xls_bytes({"S1": [["a"]]})) == "xls"


def test_sniff_workbook_format_rejects_unknown():
    with pytest.raises(ValueError):
        sniff_workbook_format(b"not a workbook")


def test_market_value_multiplier_recognizes_units():
    assert market_value_multiplier("Market Value") == Decimal("1")
    assert market_value_multiplier("Market Value (Rs.)") == Decimal("1")
    assert market_value_multiplier("Market Value (₹)") == Decimal("1")
    assert market_value_multiplier("Market Value (Rs. in Lakhs)") == Decimal("100000")
    assert market_value_multiplier("Market Value (Rs. in Crores)") == Decimal("10000000")


def test_market_value_multiplier_raises_on_unrecognized_unit():
    # A future AMC's "(Rs. in '000s)" must not be silently treated as bare rupees --
    # the parenthetical names a scale this function doesn't recognize.
    with pytest.raises(ValueError):
        market_value_multiplier("Market Value (Rs. in '000s)")

# append to backend/tests/services/analytics/test_portfolio_disclosure_parsing.py
from app.models.enums import AssetType

_HOLDING_ROWS = [
    ["Name of the Instrument", "ISIN", "Industry", "Quantity", "Market Value (Rs. in Lakhs)", "% to NAV"],
    ["Equity & Equity Related", "", "", "", "", ""],
    ["HDFC Bank Ltd", "INE040A01034", "Banks", 1000, 15.00, 2.5],
    ["Reliance Industries Ltd", "INE002A01018", "Petroleum Products", 500, 20.00, 3.3],
    ["Sub Total", "", "", "", 35.00, 5.8],
    ["Money Market Instruments", "", "", "", "", ""],
    ["Net Receivables/Payables", "", "", "", 1.00, 0.2],
    ["Total", "", "", "", 36.00, 6.0],
]


def test_find_header_row_locates_header_below_a_title_banner():
    rows = [["Nippon India Large Cap Fund - Portfolio as on 31-Aug-2026"], [], *_HOLDING_ROWS]
    result = find_header_row(rows)
    assert result is not None
    row_index, columns = result
    assert row_index == 2
    assert columns["instrument_name"] == 0
    assert columns["isin"] == 1
    assert columns["market_value"] == 4
    assert columns["pct_to_nav"] == 5


def test_find_header_row_returns_none_for_a_non_holdings_sheet():
    assert find_header_row([["Index"], ["Scheme Code", "Scheme Name"], ["YB01", "Scheme A"]]) is None


def test_find_header_row_matches_pct_to_aum_header():
    # Choice and Aditya Birla Sun Life label this column "% to AUM", not "% to NAV".
    rows = [["Name of the Instrument", "ISIN", "Industry", "Quantity", "Market Value", "% to AUM"],
            ["HDFC Bank Ltd", "INE040A01034", "Banks", 1000, 15.00, 2.5]]
    result = find_header_row(rows)
    assert result is not None
    assert result[1]["pct_to_nav"] == 5


def test_parse_sheet_rows_stops_at_total_catches_sub_total_and_stamps_asset_type():
    _, columns = find_header_row(_HOLDING_ROWS)
    holdings = parse_sheet_rows(_HOLDING_ROWS[1:], columns, Decimal("100000"))
    # 2 equities + the cash/receivables line, not the "Sub Total" or "Total" rows --
    # the original exact-match terminator check would have ingested "Sub Total" as a
    # third equity holding, double-counting that section's weight.
    assert len(holdings) == 3
    assert holdings[0].instrument_name == "HDFC Bank Ltd"
    assert holdings[0].isin == "INE040A01034"
    assert holdings[0].market_value == Decimal("1500000.00")
    assert holdings[0].asset_type == AssetType.EQUITY_DOMESTIC
    assert holdings[1].asset_type == AssetType.EQUITY_DOMESTIC
    assert holdings[2].instrument_name == "Net Receivables/Payables"
    assert holdings[2].isin is None
    assert holdings[2].asset_type == AssetType.CASH_EQUIVALENT


def test_parse_sheet_rows_defaults_to_other_before_any_section_header():
    rows = [["Some Unlabeled Instrument", "", "", "", 1.00, 0.1], ["Total", "", "", "", 1.00, 0.1]]
    _, columns = find_header_row([_HOLDING_ROWS[0], *rows])
    holdings = parse_sheet_rows(rows, columns, Decimal("1"))
    assert holdings[0].asset_type == AssetType.OTHER


def test_parse_sheet_rows_handles_percent_values_written_as_text():
    rows = [["HDFC Bank Ltd", "INE040A01034", "Banks", 1000, "15.00", "2.50%"],
            ["Total", "", "", "", "15.00", "2.50%"]]
    _, columns = find_header_row([_HOLDING_ROWS[0], *rows])
    holdings = parse_sheet_rows(rows, columns, Decimal("1"))
    assert len(holdings) == 1
    assert holdings[0].pct_to_nav == Decimal("2.50")


def test_parse_workbook_skips_non_holdings_sheets_and_normalizes_crores():
    crore_rows = [
        ["Name of the Instrument", "ISIN", "Industry", "Quantity", "Market Value (Rs. in Crores)", "% to NAV"],
        ["Equity & Equity Related", "", "", "", "", ""],
        ["Infosys Ltd", "INE009A01021", "IT - Software", 2000, 1.50, 4.0],
    ]
    data = _xlsx_bytes({"Index": [["Scheme Code", "Scheme Name"], ["YB01", "Fund A"]], "Fund A": crore_rows})
    sheets = parse_workbook(data)
    names = {s.sheet_name for s in sheets}
    assert names == {"Fund A"}  # "Index" has no SEBI header row -- silently skipped, not an error
    fund_a = next(s for s in sheets if s.sheet_name == "Fund A")
    assert fund_a.error is None
    assert fund_a.holdings[0].market_value == Decimal("15000000.00")
    assert fund_a.holdings[0].asset_type == AssetType.EQUITY_DOMESTIC

# append to backend/tests/services/analytics/test_portfolio_disclosure_parsing.py
import zipfile
import re
from functools import lru_cache
from pathlib import Path

import xlrd

from app.services.analytics.portfolio_disclosure_parsing import extract_zip_members


def test_extract_zip_members_returns_only_workbook_files():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("Fund A.xlsx", _xlsx_bytes({"Sheet1": [["a"]]}))
        zf.writestr("Fund B.xls", _xls_bytes({"Sheet1": [["a"]]}))
        zf.writestr("readme.txt", "not a workbook")
    members = extract_zip_members(buf.getvalue())
    names = {name for name, _ in members}
    assert names == {"Fund A.xlsx", "Fund B.xls"}


def test_fraction_scale_is_detected_per_sheet_and_stops_at_grand_total():
    rows = [
        _HOLDING_ROWS[0],
        ["Equity & Equity Related"],
        ["Bank", "INE040A01034", "Banks", 10, 1, "0.925"],
        ["Sub Total", "", "", "", 1, "0.925"],
        [],
        ["Net Receivables / (Payables)", "", "", "", 1, "0.075"],
        ["GRAND TOTAL", "", "", "", 2, 1],
        ["Annexure Bank", "INE040A01034", "Banks", 10, 10000, 500],
    ]
    point_rows = [list(r) for r in rows[:7]]
    point_rows[2][-1] = "92.5"
    point_rows[5][-1] = "7.5"
    point_rows[6][-1] = "100"
    sheets = parse_workbook(_xlsx_bytes({"Fraction": rows, "Points": point_rows}))
    assert len(sheets) == 2
    for sheet in sheets:
        assert len(sheet.holdings) == 2
        assert sum((h.pct_to_nav for h in sheet.holdings), Decimal(0)) == Decimal(100)
        assert sheet.holdings[-1].asset_type == AssetType.CASH_EQUIVALENT


def test_header_positions_isin_first_and_shifted_isin_cells():
    rows = [
        ["", "ISIN Code", "Name of Instrument", "Industry", "Quantity", "Mkt Value Rs. in Lacs", "% of Net Asset"],
        ["", "Equity & Equity Related"],
        ["code", "**", "Bank", "Banks", 10, 1, 100, "INE040A01034"],
    ]
    sheet = parse_workbook(_xlsx_bytes({"Holding": rows}))[0]
    assert sheet.error is None
    assert sheet.holdings[0].isin == "INE040A01034"
    assert sheet.holdings[0].market_value == Decimal(100000)


@pytest.mark.parametrize("text", ["NaN", "Infinity", "-Infinity", "#VALUE!"])
def test_nonfinite_or_error_cells_never_become_holdings(text):
    rows = [["Bank", "INE040A01034", "Banks", 10, text, 100]]
    _, columns = find_header_row([_HOLDING_ROWS[0], *rows])
    assert parse_sheet_rows(rows, columns, Decimal(1)) == []


def test_merged_name_header_and_total_label_in_another_column():
    rows = [
        ["Name of Instrument", "", "", "ISIN", "Industry", "Quantity", "Market Value (Rs.in Lacs)", "% to NAV"],
        ["Equity & Equity Related"],
        ["", "", "Bank", "INE040A01034", "Banks", 10, 1, 0.95],
        ["", "", "Net Current Assets", "", "", "", 1, 0.05],
        ["", "", "", "", "Grand Total", "", 2, 1],
    ]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert [h.instrument_name for h in sheet.holdings] == ["Bank", "Net Current Assets"]
    assert sheet.holdings[0].pct_to_nav == Decimal(95)


def test_numeric_category_headers_are_not_positions_and_underscore_total_stops():
    rows = [
        _HOLDING_ROWS[0],
        ["Debt Instruments", "", "", "", 10, 0.95],
        ["Listed / Awaiting Listing On Stock Exchanges", "", "", "", 10, 0.95],
        ["Government Securities", "", "", "", 10, 0.95],
        ["Bond", "IN0020250018", "Sovereign", 10, 10, 0.95],
        ["Net Receivables", "", "", "", 1, 0.05],
        ["GRAND_TOTAL", "", "", "", 11, 1],
        ["Notes:", "", "", "", "", ""],
        ["Bond annexure", "IN0020250018", "Sovereign", 10, 10, 0.95],
    ]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert len(sheet.holdings) == 2
    assert sheet.holdings[0].asset_type == AssetType.DEBT
    assert sum((h.pct_to_nav for h in sheet.holdings), Decimal(0)) == Decimal(100)


def test_declared_total_is_scaled_but_derivative_positions_remain_disclosed():
    rows = [_HOLDING_ROWS[0], ["Equity & Equity Related"],
            ["Bank", "INE040A01034", "Banks", 10, 1, 0.9],
            ["Derivatives"], ["Bank Future", "", "Banks", -10, -1, -0.8],
            ["Net Current Assets", "", "", "", 1, 0.1],
            ["Grand Total", "", "", "", 2, 1]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert sheet.declared_total == Decimal(100)
    assert sheet.holdings[1].pct_to_nav == Decimal(-80)
    assert sheet.holdings[1].asset_type is AssetType.DERIVATIVE
    assert sum((h.pct_to_nav for h in sheet.holdings), Decimal(0)) == Decimal(20)


def test_sector_raw_is_not_normalized_and_cash_liabilities_override_previous_section():
    rows = [_HOLDING_ROWS[0], ["Equity & Equity Related"],
            ["Bank", "INE040A01034", "  Banks /\n Rating  ", 10, 1, 95],
            ["Foreign Securities"], ["NET CURRENT LIABILITIES", "", "", "", 1, 5]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert sheet.holdings[0].sector_raw == "  Banks /\n Rating  "
    assert sheet.holdings[1].asset_type is AssetType.CASH_EQUIVALENT


def test_scheme_marker_blocks_do_not_mix_portfolios_or_annexures():
    rows = []
    for code in ("001", "002"):
        rows += [[f"SCHEME CODE{code}STARTS"], [f"SCHEME: Fund {code}"], _HOLDING_ROWS[0],
                 ["Equity & Equity Related"], ["Bank", "INE040A01034", "Banks", 1, 1, 100],
                 [f"TOTAL : Fund {code}", "", "", "", 1, ""],
                 ["Defaulted Annexure", "INE040A01034", "Banks", 1, 1, 500],
                 [f"SCHEME CODE{code}ENDS"]]
    sheets = parse_workbook(_xlsx_bytes({"EXPOSURE": rows}))
    assert [s.sheet_name for s in sheets] == ["EXPOSURE:001", "EXPOSURE:002"]
    assert all(len(s.holdings) == 1 for s in sheets)
    assert all(s.declared_total is None for s in sheets)


def test_irs_section_is_derivative():
    rows = [_HOLDING_ROWS[0], ["IRS"], ["IRS - Contract", "", "", 1, 1, 5]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert sheet.holdings[0].asset_type is AssetType.DERIVATIVE


def test_zero_value_segregated_positions_with_tiny_weight_marker_are_retained():
    rows = [_HOLDING_ROWS[0], ["Debt Instruments"],
            ["Defaulted Bond", "INE528G08352", "ICRA-D", 950, 0, "*"],
            ["Net Current Assets", "", "", "", 0, "*"]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert len(sheet.holdings) == 2
    assert all(h.market_value == Decimal(0) and h.pct_to_nav == Decimal(0) for h in sheet.holdings)


def test_gold_silver_related_section_classifies_its_positions():
    rows = [_HOLDING_ROWS[0], ["Gold/Silver related"], ["Metal Position", "", "", 10, 1, 100]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert sheet.holdings[0].asset_type is AssetType.GOLD_SILVER


def test_invalid_declared_total_is_a_sheet_error_and_does_not_lose_other_sheets():
    good = [_HOLDING_ROWS[0], ["Bank", "INE040A01034", "Banks", 1, 1, 100]]
    bad = [*good, ["Grand Total", "", "", "", 1, "#VALUE!"]]
    sheets = {s.sheet_name: s for s in parse_workbook(_xlsx_bytes({"Good": good, "Bad": bad}))}
    assert sheets["Good"].error is None
    assert sheets["Good"].holdings
    assert sheets["Bad"].error is not None
    assert "declared total" in sheets["Bad"].error


_FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "portfolio_disclosures"
_NON_HOLDINGS = {"index", "scheme", "sheet1-index", "derivative", "annexure", "f & o disclosure",
                 "common notes", "xdo_metadata"}
# Only legitimately exceptional totals may be listed here, with a reason per sheet.
REAL_SHEET_SUM_ALLOWLIST = {
    ('quant-all-schemes_asof_2026-09-30.xlsx', 'qMES'):
        ('Source-duplicated TREPS/NCA rows inflate the non-derivative sum.', Decimal('115.744795'), Decimal('100.000000')),
    ('quant-all-schemes_asof_2026-09-30.xlsx', 'qMAB'):
        ('Source-duplicated TREPS/NCA rows inflate the non-derivative sum.', Decimal('102.254827'), Decimal('100.000000')),
    ('quant-arbitrage_asof_2026-09-30.xlsx', 'quant_Arbitrage_Fund'):
        ('Source-duplicated TREPS/NCA rows inflate the non-derivative sum.', Decimal('102.254827'), Decimal('100.000000')),
    ('tata_asof_2026-09-30.xlsx', 'TMAOF'):
        ('Negative commodity hedges are memo; non-derivative rows total 94.73 against a 100 Grand Total.', Decimal('94.730000'), Decimal('100.000000')),
    ('uti-consolidated_asof_2026-09-30.zip', 'EXPOSURE:176'):
        ('Source numeric weights total 101.28%; NCA has a rounded-zero marker.', Decimal('101.280000'), None),
    ('uti-consolidated_asof_2026-09-30.zip', 'EXPOSURE:747'):
        ('Zero-NAV segregated YES Bank portfolio.', Decimal('0.000000'), None),
    ('uti-consolidated_asof_2026-09-30.zip', 'EXPOSURE:748'):
        ('Zero-NAV segregated YES Bank portfolio.', Decimal('0.000000'), None),
}


@lru_cache(maxsize=None)
def _fixture_workbooks(filename):
    data = (_FIXTURES / filename).read_bytes()
    if data.startswith(b"PK"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "[Content_Types].xml" not in archive.namelist():
                return [(n, archive.read(n)) for n in archive.namelist() if n.lower().endswith((".xls", ".xlsx"))]
    return [(filename, data)]


def _raw_sheets(data):
    if data.startswith(b"PK"):
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        try:
            for ws in wb:
                if ws.title.lower() in _NON_HOLDINGS or ws.title.lower().startswith("derivative"):
                    continue
                rows = []
                for row in ws.iter_rows(values_only=True):
                    rows.append(list(row))
                    if any(re.match(r"grand\s+total", str(c or ""), re.I) for c in row):
                        break
                yield ws.title, rows
        finally:
            wb.close()
    else:
        wb = xlrd.open_workbook(file_contents=data)
        try:
            for ws in wb.sheets():
                yield ws.name, [ws.row_values(i) for i in range(ws.nrows)]
        finally:
            wb.release_resources()


def _fixture_cases():
    cases = []
    for path in sorted(_FIXTURES.iterdir()):
        if path.suffix.lower() not in {".xls", ".xlsx", ".zip"}:
            continue
        for member, data in _fixture_workbooks(path.name):
            for name, rows in _raw_sheets(data):
                if name.lower() in _NON_HOLDINGS or "performance" in name.lower():
                    continue
                # Independent SEBI-header inventory, so parser skips cannot hide a fixture.
                is_holdings = any(
                    re.search(r"name.*instrument|instrument.*name|company/issuer", " ".join(str(c or "") for c in row), re.I)
                    and any("%" in str(c) and re.search(r"nav|aum|asset", str(c), re.I) for c in row)
                    for row in rows[:40]
                )
                if is_holdings:
                    codes = [re.search(r"SCHEME CODE(\d+)STARTS", str(c)) for row in rows for c in row]
                    codes = [m.group(1) for m in codes if m]
                    if codes:
                        cases.extend((path.name, member, f"{name}:{code}") for code in codes)
                    else:
                        cases.append((path.name, member, name))
    return cases


@lru_cache(maxsize=None)
def _parsed_fixture(filename, member):
    data = dict(_fixture_workbooks(filename))[member]
    return {s.sheet_name: s for s in parse_workbook(data)}


@pytest.mark.parametrize("filename,member,name", _fixture_cases())
def test_every_real_holdings_sheet_parses_with_a_plausible_scaled_total(filename, member, name, request):
    parsed = _parsed_fixture(filename, member)
    assert name in parsed, f"Holdings sheet skipped: {filename}/{member}/{name}"
    sheet = parsed[name]
    assert sheet.error is None, sheet.error
    assert sheet.holdings, f"Empty holdings: {filename}/{name}"
    declared = sheet.declared_total
    all_sum = sum((h.pct_to_nav for h in sheet.holdings), Decimal(0))
    nd_sum = sum((h.pct_to_nav for h in sheet.holdings if h.asset_type != AssetType.DERIVATIVE), Decimal(0))
    # A Grand Total reconciles on either basis (derivatives are in NAV for some AMCs, memo for others).
    total = all_sum if declared is None else (all_sum if abs(all_sum - declared) <= abs(nd_sum - declared) else nd_sum)
    key = (filename, name)
    if key in REAL_SHEET_SUM_ALLOWLIST:
        reason, expected_sum, expected_declared = REAL_SHEET_SUM_ALLOWLIST[key]
        assert reason
        assert total.quantize(Decimal("0.000001")) == expected_sum.quantize(Decimal("0.000001")), f"{key}: {total}"
        assert (declared.quantize(Decimal("0.000001")) if declared is not None else None) == expected_declared
        # An exception to the fixture gate never becomes an exception to Write.
        import uuid
        from datetime import date
        from app.models.reference import Scheme, DisclosureBatch, SchemeHolding
        from app.models.enums import DisclosureBatchStatus
        from app.services.analytics.amfi_factsheet_client import build_families
        from app.services.analytics.portfolio_disclosure_writer import write_disclosure_batch
        db = request.getfixturevalue("db_session")
        scheme = Scheme(id=uuid.uuid4(), name=name, base_name=name, amc_name="Fixture AMC", sebi_category="Test")
        db.add(scheme)
        db.flush()
        write_disclosure_batch(db, build_families([scheme])[0], date(2026, 9, 30), "Fixture AMC",
                               "https://example.test/fixture", sheet.holdings, "EXACT", Decimal(1),
                               declared_total=declared)
        batch = db.query(DisclosureBatch).filter_by(scheme_id=scheme.id).one()
        assert batch.status is DisclosureBatchStatus.FAILED
        assert db.query(SchemeHolding).filter_by(batch_id=batch.id).count() == len(sheet.holdings)
    elif declared is not None:
        assert Decimal(90) <= declared <= Decimal(101), f"{key}: declared={declared}"
        assert abs(total - declared) <= Decimal("1.5"), f"{key}: sum={total}, declared={declared}"
    else:
        assert Decimal(90) <= total <= Decimal(101), f"{key}: sum={total}"


# FIX ROUND 1: regressions from the independent review.
def test_fix_h1_numeric_position_named_like_a_section_is_retained():
    rows = [_HOLDING_ROWS[0], ["Debt Instruments"],
            ["Government Securities", "IN0020250018", "Sovereign", 10, 28597, 3.56],
            ["Government Securities", "", "", 10, 5, 0.5]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert len(sheet.holdings) == 2
    assert sheet.holdings[0].isin == "IN0020250018"


def test_fix_h1_icici_banpsu_keeps_every_government_security_position():
    filename = "icici-prudential_asof_2026-09-30-trimmed.zip"
    member = "ICICI Prudential Banking and PSU Debt Fund.xlsx"
    sheet = _parsed_fixture(filename, member)["BANPSU"]
    rows = dict(_raw_sheets(dict(_fixture_workbooks(filename))[member]))["BANPSU"]
    header, columns = find_header_row(rows)
    source = [r for r in rows[header+1:] if str(r[columns["instrument_name"]]).strip() == "Government Securities"
              and re.fullmatch(r"IN[A-Z0-9]{10}", str(r[columns["isin"]]))]
    positions = [h for h in sheet.holdings if h.instrument_name == "Government Securities" and h.isin]
    assert len(source) >= 9
    assert len(positions) == len(source)
    assert any(h.isin == "IN0020250018" for h in positions)


@pytest.mark.parametrize("marker", ["^", "$", "*"])
def test_fix_m1_tiny_weight_markers_keep_nonzero_positions(marker):
    rows = [_HOLDING_ROWS[0], ["Equity & Equity Related"],
            ["Oil India", "INE274J01014", "Oil", 10, 157, marker]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert len(sheet.holdings) == 1
    assert sheet.holdings[0].market_value == Decimal(15700000)
    assert sheet.holdings[0].pct_to_nav == Decimal(0)


def test_fix_m2_fraction_rows_with_explicit_percent_total_are_scaled_from_rows():
    rows = [_HOLDING_ROWS[0], ["Equity & Equity Related"],
            ["Bank", "INE040A01034", "Banks", 10, 1, 0.95],
            ["Net Current Assets", "", "", "", 1, 0.05],
            ["Grand Total", "", "", "", 2, "100%"]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert sheet.error is None
    assert sheet.declared_total == Decimal(100)
    assert sum((h.pct_to_nav for h in sheet.holdings), Decimal(0)) == Decimal(100)


def test_fix_m2_inconsistent_numeric_total_scale_is_rejected():
    rows = [_HOLDING_ROWS[0], ["Bank", "INE040A01034", "Banks", 10, 1, 0.95],
            ["Net Current Assets", "", "", "", 1, 0.05],
            ["Grand Total", "", "", "", 2, 100]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert sheet.error is not None and "scale" in sheet.error.lower()
    assert sheet.holdings  # retained for FAILED-batch diagnostics


def test_fix_h3_arbitrage_keeps_long_equity_and_types_short_or_contract_rows():
    rows = [_HOLDING_ROWS[0], ["ARBITRAGE"],
            ["Tata Steel Limited", "INE081A01020", "Metals", 10, 1, 70],
            ["Tata Steel Limited Oct26", "", "Stock Futures", -10, -1, -60],
            ["Bank LTD^", "INE040A01034", "Banks", -10, -1, -5],
            ["Bank Oct2026", "", "Banks", 10, 1, 2],
            ["Index contract", "", "Index Options", 1, 1, 1]]
    holdings = parse_workbook(_xlsx_bytes({"Fund": rows}))[0].holdings
    assert holdings[0].asset_type is AssetType.EQUITY_DOMESTIC
    assert all(h.asset_type is AssetType.DERIVATIVE for h in holdings[1:])


@pytest.mark.parametrize("filename,member,name", [
    ("dsp_asof_2026-09-30.zip", "DSP Equity FOF ISIN Portfolio as on 30 Sep 2026.xlsx", "ARBITRAGE"),
    ("dsp_asof_2026-09-30.zip", "DSP Equity FOF ISIN Portfolio as on 30 Sep 2026.xlsx", "DAAF"),
    ("dsp_asof_2026-09-30.zip", "DSP Equity FOF ISIN Portfolio as on 30 Sep 2026.xlsx", "VALUE"),
    ("tata_asof_2026-09-30.xlsx", "tata_asof_2026-09-30.xlsx", "TAF"),
    ("tata_asof_2026-09-30.xlsx", "tata_asof_2026-09-30.xlsx", "TMAOF"),
])
def test_fix_h3_real_negative_equity_exposures_are_derivatives(filename, member, name):
    sheet = _parsed_fixture(filename, member)[name]
    negatives = [h for h in sheet.holdings if h.quantity is not None and h.quantity < 0]
    assert negatives
    assert all(h.asset_type is AssetType.DERIVATIVE for h in negatives)


def test_fix_h3_physical_gold_silver_do_not_inherit_equity():
    rows = [_HOLDING_ROWS[0], ["Equity & Equity Related"], ["COMMODITIES & COMMODITIES RELATED"],
            ["A) Physical Commodities"], ["SILVER (30 KG) COMMODITY", "", "", 10, 1, 100]]
    assert parse_workbook(_xlsx_bytes({"Fund": rows}))[0].holdings[0].asset_type is AssetType.GOLD_SILVER


def test_fix_l1_biff_skips_derivative_and_performance_sheets():
    data = _xls_bytes({"Holdings": _HOLDING_ROWS, "DerivativeMemo": _HOLDING_ROWS,
                       "Fund Performance": _HOLDING_ROWS})
    assert [s.sheet_name for s in parse_workbook(data)] == ["Holdings"]


def test_fix_l2_terminal_banner_before_header_does_not_truncate_xlsx():
    rows = [["Net Assets", "100"], *_HOLDING_ROWS]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert len(sheet.holdings) == 3


@pytest.mark.parametrize("fmt", ["xlsx", "xls"])
def test_fix_m2_percent_formatted_gross_position_over_one_with_negative_cash(fmt):
    rows = [_HOLDING_ROWS[0], ["Mutual Fund Units"],
            ["Gold ETF", "INF204KB17I5", "", 10, 1, 1.02],
            ["Net Current Assets", "", "", "", -1, -0.02],
            ["Grand Total", "", "", "", 1, 1]]
    if fmt == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in rows: ws.append(row)
        for row in (3, 4, 5): ws.cell(row, 6).number_format = "0.00%"
    else:
        wb = xlwt.Workbook()
        ws = wb.add_sheet("Fund")
        style = xlwt.easyxf(num_format_str="0.00%")
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                ws.write(r, c, cell, style if r >= 2 and c == 5 else xlwt.Style.default_style)
    out = io.BytesIO()
    wb.save(out)
    sheet = parse_workbook(out.getvalue())[0]
    assert sheet.error is None
    assert sheet.holdings[0].pct_to_nav == Decimal(102)
    assert sheet.holdings[1].pct_to_nav == Decimal(-2)
    assert sheet.declared_total == Decimal(100)


def test_fix_h1_position_named_notes_does_not_end_a_live_isin_table():
    rows = [_HOLDING_ROWS[0], ["Debt Instruments"],
            ["Notes:", "IN000438C043", "Sovereign", 10, 1, 40],
            ["Bond", "IN000439C041", "Sovereign", 10, 1, 60],
            ["Grand Total", "", "", "", 2, 100]]
    sheet = parse_workbook(_xlsx_bytes({"Fund": rows}))[0]
    assert len(sheet.holdings) == 2


@pytest.mark.parametrize("name,member", [
    ("BAF", "ICICI Prudential Balanced Advantage Fund.xlsx"),
    ("ARBITRAGE", "ICICI Prudential Arbitrage Fund.xlsx"),
])
def test_fix_h1_icici_numeric_aggregate_summaries_do_not_duplicate_positions(name, member):
    sheet = _parsed_fixture("icici-prudential_asof_2026-09-30-trimmed.zip", member)[name]
    assert abs(sum((h.pct_to_nav for h in sheet.holdings), Decimal(0))-sheet.declared_total) <= Decimal("1.5")


def test_fix_h1_absl_notes_named_bond_does_not_drop_tail_positions():
    sheet = _parsed_fixture("aditya-birla-sun-life_asof_2026-09-30.zip", "30092026_ABSLMF_MonthlyDisclosure.xls")["BSLIF"]
    assert any(h.isin == "IN000438C043" for h in sheet.holdings)
    assert abs(sum((h.pct_to_nav for h in sheet.holdings), Decimal(0))-sheet.declared_total) <= Decimal("1.5")
